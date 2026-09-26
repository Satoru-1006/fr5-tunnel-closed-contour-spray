from copy import deepcopy
import re
from pathlib import Path

from docx import Document
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Pt, RGBColor


SRC = Path(r"C:\Users\86198\Desktop\report.docx")
OUT = Path(r"C:\Users\86198\Desktop\report_论文规范修订版.docx")


def set_run_font(run, east="宋体", west="Times New Roman", size=12, bold=None):
    run.font.name = west
    run._element.get_or_add_rPr().get_or_add_rFonts().set(qn("w:eastAsia"), east)
    run._element.rPr.rFonts.set(qn("w:ascii"), west)
    run._element.rPr.rFonts.set(qn("w:hAnsi"), west)
    run.font.size = Pt(size)
    run.font.color.rgb = RGBColor(0, 0, 0)
    if bold is not None:
        run.bold = bold


def replace_paragraph_text(p, text, size=12, bold=None):
    for r in list(p.runs):
        p._p.remove(r._r)
    r = p.add_run(text)
    set_run_font(r, size=size, bold=bold)
    return p


def insert_before(anchor, text, style="Normal", bold=False):
    p = OxmlElement("w:p")
    anchor._p.addprevious(p)
    newp = anchor._parent.add_paragraph()
    newp._p.getparent().remove(newp._p)
    p.addnext(newp._p)
    newp.style = style
    r = newp.add_run(text)
    set_run_font(r, size=12, bold=bold)
    newp.paragraph_format.first_line_indent = Pt(24) if style == "Normal" else None
    newp.paragraph_format.space_after = Pt(6)
    newp.paragraph_format.line_spacing = 1.5
    return newp


def insert_after(anchor, text, style="Normal", bold=False):
    newp = anchor._parent.add_paragraph()
    newp._p.getparent().remove(newp._p)
    anchor._p.addnext(newp._p)
    newp.style = style
    r = newp.add_run(text)
    set_run_font(r, size=12, bold=bold)
    newp.paragraph_format.first_line_indent = Pt(24) if style == "Normal" else None
    newp.paragraph_format.space_after = Pt(6)
    newp.paragraph_format.line_spacing = 1.5
    return newp


def find_para(doc, starts):
    for p in doc.paragraphs:
        if p.text.strip().startswith(starts):
            return p
    raise ValueError(starts)


def find_heading(doc, starts, style_name=None):
    for p in doc.paragraphs:
        if p.text.strip().startswith(starts) and (style_name is None or p.style.name == style_name):
            return p
    raise ValueError((starts, style_name))


def remove_between(start, end):
    node = start._p.getnext()
    while node is not None and node is not end._p:
        nxt = node.getnext()
        node.getparent().remove(node)
        node = nxt


def strip_all_shading(doc):
    root = doc._element
    for shd in root.xpath(".//w:shd"):
        shd.set(qn("w:val"), "clear")
        shd.set(qn("w:color"), "auto")
        shd.set(qn("w:fill"), "FFFFFF")


def force_black_fonts(doc):
    for p in doc.paragraphs:
        for r in p.runs:
            r.font.color.rgb = RGBColor(0, 0, 0)
            rPr = r._element.get_or_add_rPr()
            color = rPr.find(qn("w:color"))
            if color is None:
                color = OxmlElement("w:color")
                rPr.append(color)
            color.set(qn("w:val"), "000000")
            color.attrib.pop(qn("w:themeColor"), None)
    for table in doc.tables:
        for row in table.rows:
            for cell in row.cells:
                for p in cell.paragraphs:
                    for r in p.runs:
                        r.font.color.rgb = RGBColor(0, 0, 0)


def normalize_captions(doc):
    chapter = 0
    fig = {}
    tab = {}
    for p in doc.paragraphs:
        t = re.sub(r"\s+", "", p.text)
        m = re.match(r"第(\d+)章", t)
        if m:
            chapter = int(m.group(1))
        if re.match(r"^图\s*\d+", p.text.strip()):
            fig[chapter] = fig.get(chapter, 0) + 1
            replace_paragraph_text(p, f"图{chapter}-{fig[chapter]}", size=9)
            p.alignment = WD_ALIGN_PARAGRAPH.CENTER
            p.paragraph_format.keep_with_next = False
            p.paragraph_format.space_before = Pt(3)
            p.paragraph_format.space_after = Pt(6)
        elif re.match(r"^表\s*\d+", p.text.strip()):
            tab[chapter] = tab.get(chapter, 0) + 1
            replace_paragraph_text(p, f"表{chapter}-{tab[chapter]}", size=9)
            p.alignment = WD_ALIGN_PARAGRAPH.CENTER
            p.paragraph_format.keep_with_next = True
            p.paragraph_format.space_before = Pt(6)
            p.paragraph_format.space_after = Pt(3)


def add_version_workflows(doc):
    workflows = {
        "1.2 第一版": (
            "本版性质界定：第一版属于机器人运动仿真的几何与路径基线，只处理隧道截面、喷涂轨迹点、喷距和路径连续性，不求解空气流场，也不包含液滴、颗粒扰动或沉积过程。流体仿真属于后文单独算例，不能用其颗粒图替代本版运动学结果。",
            "本版运动仿真按以下顺序实施：第一步，根据马蹄形隧道截面尺寸建立二维轮廓并离散为有序壁面点；第二步，依据喷涂间距和闭合方向生成连续轮廓路径；第三步，将二维点扩展到机器人工作坐标系并检查点序、间距和首尾连接；第四步，输出路径CSV与三维预览图；第五步，以几何闭合性、路径覆盖性和相邻点距离作为验收指标。本版没有独立流体计算步骤，其成果仅作为后续机器人求解的目标轨迹输入。"),
        "1.3 第二版": (
            "本版性质界定：第二版仍是纯机器人算法层面的运动仿真，重点是用Python建立简化六轴机构、读取目标路径并形成可视化动画。动画中的机械臂与轨迹线只表达关节运动和末端位姿，不代表喷雾颗粒、空气扰动或材料沉积。",
            "本版运动仿真按以下顺序实施：第一步，读取第一版的有序路径点并统一坐标与尺度；第二步，建立简化连杆和关节坐标变换；第三步，对每个目标点计算末端位置和姿态目标；第四步，生成相邻帧之间的关节状态并绘制机械臂动画；第五步，从路径跟随、姿态变化和动画连贯性三个方面复核结果。该阶段尚未接入FR5官方模型、MoveIt2碰撞环境或流体求解器，因此结论仅限于算法原型可运行。"),
        "1.4 第三版": (
            "本版性质界定：第三版是机器人模型和工具坐标系的运动仿真升级，核心是由简化机构切换到FAIRINO FR5 V6模型，并引入150 mm虚拟喷头TCP。该TCP是运动学设计参数，不是液滴喷射模型；本版仍不计算任何连续相或颗粒相。",
            "本版运动仿真按以下顺序实施：第一步，导入并核对FR5 V6关节轴、连杆尺寸和关节限位；第二步，在法兰坐标系定义150 mm工具偏置；第三步，将原末端目标换算为法兰目标位姿；第四步，执行逆运动学并检查关节是否落在允许范围；第五步，通过正运动学回代比较TCP目标与实际位置。最终输出模型、TCP元数据和关节轨迹，流体仿真则在独立章节使用自身的喷嘴入口条件。"),
        "1.5 第四版": (
            "本版性质界定：第四版属于ROS2与MoveIt2环境中的严格运动规划仿真，研究对象是FR5关节轨迹、TCP位姿、正运动学误差和碰撞状态。碰撞对象是机器人与几何环境，不是喷雾粒子之间的相互作用。",
            "本版运动仿真按以下顺序实施：第一步，将路径点、四元数和TCP参数转换为MoveIt2输入；第二步，加载FR5模型、规划组和简化隧道碰撞体；第三步，对目标点求解机器人状态并生成关节轨迹；第四步，利用正运动学复算TCP位置和朝向误差；第五步，逐状态检查自碰撞与环境碰撞；第六步，导出关节轨迹和严格验证报告。只有这些机器人算法指标通过，才能进入后续平滑与工艺段设计。"),
        "1.6 第五版": (
            "本版性质界定：第五版关注运动轨迹的时间参数化和动力学质量门禁，jerk表示关节加速度对时间的变化率，与流体湍动或粒子扰动没有关系。文中出现的尖峰只用于判断机器人轨迹是否平滑。",
            "本版运动仿真按以下顺序实施：第一步，读取MoveIt2关节路径及关节限制；第二步，使用Ruckig分配时间并生成速度、加速度和jerk序列；第三步，按关节上限计算归一化比值；第四步，区分真实超限、离散差分造成的尖峰和阈值附近的告警；第五步，将平滑前后结果与TCP速度复算结果对照；第六步，输出质量报告。本版不运行流体求解，动力学曲线也不能解释为喷雾流场曲线。"),
        "1.7 第六版": (
            "本版性质界定：第六版针对机器人逆运动学分支和几何碰撞根因开展诊断，所谓“局部窗口探针”是对关节解与碰撞状态的局部搜索，不是对流体网格或颗粒云的扰动分析。",
            "本版运动仿真按以下顺序实施：第一步，定位关节步长突增或碰撞首次出现的轨迹索引；第二步，围绕该索引建立局部状态窗口；第三步，分别改变基座位置、TCP方向、工具滚转角和IK种子；第四步，对每组候选重新执行正运动学、关节限位和碰撞检查；第五步，用根因矩阵记录成功率与失败原因；第六步，保留能够解释问题的最小修正。本版产出的是机器人构型诊断证据，不涉及流体参数。"),
        "1.8 第七版": (
            "本版性质界定：第七版把无法在整条闭合轨迹上消除的逆运动学分支切换转化为“喷涂段+停喷换姿段”的机器人作业逻辑。停喷只是一项工艺状态标记，并不意味着本版计算了喷雾粒子的启停响应。",
            "本版运动仿真按以下顺序实施：第一步，识别连续关节轨迹中的大步长位置；第二步，在不破坏有效喷涂覆盖的前提下确定分段边界；第三步，将正常路径划分为13个喷涂段；第四步，在相邻喷涂段之间插入12个停喷换姿段；第五步，对喷涂段和换姿段分别检查最大关节步长；第六步，抽查插值状态的碰撞与限位；第七步，形成可审计的段落清单。流体仿真如需使用启停信息，只能在独立模块中只读该清单。"),
        "1.9 第八版": (
            "本版性质界定：第八版是开放拱形路径和多视图表达的运动仿真版本，显示内容包括机器人姿态、轨迹、隧道几何和观察视角。即使动画用于说明喷涂任务，也不包含空气流场、颗粒云、液滴破碎或沉积膜厚。",
            "本版运动仿真按以下顺序实施：第一步，从闭合轮廓中删除不参与喷涂的底部横线；第二步，根据开放拱形壁面法向重建TCP朝向；第三步，通过基座后移和滚转自由度调整获得较舒展构型；第四步，生成181点连续逆运动学结果；第五步，输出正视、侧视、俯视和三维动画；第六步，再次执行MoveIt2、Ruckig、FK和碰撞验证。多视图只用于增强机器人运动结果的可解释性。"),
    }
    for prefix, (boundary, steps) in workflows.items():
        h = find_heading(doc, prefix, "Heading 2")
        p1 = insert_after(h, boundary)
        insert_after(p1, steps)


FORMULA_TEXTS = {
    "5.1": "本节公式用于把隧道壁面几何转换为喷嘴姿态约束。选择单位法向而不是直接使用两点差值，是因为喷距会改变向量长度，却不应改变喷射方向；归一化后，路径各点都能使用同一姿态判据。计算时先用壁面点减去TCP点得到方向向量，再除以其二范数，随后把所得单位向量与四元数旋转后的工具轴做夹角比较。以报告中的最大对齐误差8.54×10⁻⁷°代入，明显小于工程校验阈值0.1°，因此姿态构造满足朝墙要求。该公式只验证机器人末端朝向，不涉及流体速度方向或颗粒扰动。",
    "5.2": "本节公式用于在喷嘴TCP目标和机器人法兰目标之间建立严格的坐标变换。逆运动学求解器控制的是法兰坐标系，如果直接把喷嘴端点当作法兰目标，150 mm工具长度会造成系统性位置偏差，因此必须使用齐次变换的逆矩阵消除工具偏置。代入当前工具参数时，法兰到TCP沿工具Z轴平移0.150 m，反求法兰位姿即沿同一轴反向平移0.150 m；例如TCP目标距壁面0.150 m时，法兰位置还需再退让0.150 m。该计算说明虚拟TCP如何进入运动学链，同时也提示实体喷枪安装后必须以标定矩阵替换。",
    "5.3": "本节公式用于在多个可行逆解中选择与上一轨迹点最连续的关节构型。仅最小化位置误差可能使求解器在等价分支间跳变，仅强调姿态又可能牺牲轨迹跟随，因此目标函数同时包含位置、姿态和相邻关节差三项。将实际权重代入后，代价函数可写为1.00×位置误差+0.40×姿态误差+0.004×连续性误差。求解时以前一点关节角作为种子，逐点比较候选解并保留总代价最小者；181个目标点均获得解，说明该权重能完成开放拱形路径求解，但仍需结合关节步长门限排除分支突跳。",
    "5.4": "本节公式用于把离散路径点转换为具有物理意义的弧长坐标，避免简单点序号掩盖局部点距差异。计算时对相邻TCP位置作差并取欧氏范数，得到每一小段长度，再逐段累加形成累计弧长；速度和时间参数化随后都基于该弧长展开。开放拱形轨迹共有181个点，因此有效相邻区间为181−1=180段。若第i段长度为Δs_i、给定TCP速度为v_i，则其初始时间估计为Δt_i=Δs_i/v_i，各段时间再交由Ruckig满足关节限制。该流程使速度、加速度和jerk检查对应真实空间行程，而不是对应任意数组索引。",
    "5.5": "本节公式用于把关节连续性、速度、加速度、jerk以及碰撞结果统一为可审查的运动仿真门限。对每个关节分别计算相邻角度差和各阶时间导数，再与模型上限比较；碰撞检查则对同一时刻的完整机器人状态求布尔结果。将交付状态的最大喷涂段关节步长3.614°代入20°门限，可得3.614/20=0.1807，即只占门限的18.07%；同时181个Ruckig后状态的自碰撞数与环境碰撞数均为0。二者联合说明当前轨迹在简化模型中连续且无碰撞，但该结论仍依赖虚拟TCP和当前碰撞几何。",
    "5.6": "本节公式用于描述流体仿真中离散颗粒相对连续空气相的运动响应，与前述机器人关节运动公式属于独立模型。采用拉格朗日方法的原因是它能够逐parcel跟踪位置、速度和边界事件，并通过Schiller–Naumann关系把相对速度转换为阻力系数。计算流程为：先由空气速度与颗粒速度之差求相对速度，再结合空气密度、动力黏度和颗粒直径计算雷诺数，随后分段求阻力系数，最后把阻力代入牛顿第二定律推进下一时刻速度。以空气密度1.2 kg/m³、运动黏度1×10⁻⁵ m²/s为已知条件时，动力黏度为1.2×10⁻⁵ Pa·s；但逐颗粒相对速度未完整留存，因此不虚构最终阻力数值。",
    "5.7": "本节公式用于解释入口粒径设定与壁面到达粒径统计之间的关系。Rosin–Rammler分布适合表达喷雾颗粒的非单一粒径特征，特征直径和分布指数可以直接控制累计分布形状。将特征直径0.09 mm、指数2和累计概率0.5代入反函数，可得入口理论中位粒径d50=0.09×[−ln(1−0.5)]^(1/2)=0.0749 mm，即74.9 μm。壁面统计D50为87.00 μm，比入口理论值高12.1 μm，增幅约16.2%。该差异反映输运和壁面筛选后的条件分布，不表示公式失效，也不能脱离parcel权重解释为质量中位粒径。",
    "5.8": "本节公式用于把目标壁面命中、开放边界逃逸和近壁空间离散程度转换为可比较指标。选择事件比例而不是直接称为沉积率，是因为日志保存的是边界事件次数，没有逐parcel质量权重；选择变异系数则是为了用标准差与均值的比值消除绝对量纲。代入6108次目标壁面事件和2679次开放边界事件，总事件数为6108+2679=8787，命中比例为6108/8787×100%=69.51%，逃逸比例为30.49%。近壁代理分布的变异系数为0.581，峰值与非零均值之比为3.99，说明相对分布不均匀，但不能换算为真实膜厚。",
}


def add_formula_explanations(doc):
    for prefix, text in FORMULA_TEXTS.items():
        h = find_heading(doc, prefix, "Heading 2")
        insert_after(h, text)


def add_ai_chapter(doc):
    refs = next(p for p in doc.paragraphs if p.style.name == "Heading 1" and p.text.strip() == "参考文献")
    h = insert_before(refs, "第7章  AI辅助使用声明", style="Heading 1")
    # Heading 1 already carries the document's chapter-break behavior; a second
    # explicit page break would create a fully blank page.
    h.paragraph_format.page_break_before = False
    p = insert_before(refs, "7.1 使用工具、用途与占比", style="Heading 2")
    insert_before(refs, "本文在撰写与整理过程中使用了DeepSeek和ChatGPT进行辅助。DeepSeek主要用于前期术语核对与思路梳理，估算占全文工作的5%；ChatGPT主要用于文章格式编排、章节结构调整、标题层级统一、公式说明扩写以及部分文字表达的规范化，估算占全文工作的15%。两项AI辅助合计约占论文整体工作的20%，该比例是对辅助性工作的估算，不代表仿真数据或技术结论由AI生成。")
    insert_before(refs, "7.2 作者核验与责任说明", style="Heading 2")
    insert_before(refs, "论文中的机器人运动仿真、OpenFOAM流体仿真、输入参数、程序文件、计算结果、图表数据、技术路线和最终结论均由作者结合项目实际情况进行核对、修改与确认。AI未替代作者完成实验、数值求解或结果验收，也未被用作原始数据来源。对于AI辅助形成的文字，作者逐项检查其与源代码、CSV/JSON报告、图形和版本记录的一致性；论文最终内容、学术规范及相关责任均由作者本人承担。")


def clean_cover(doc):
    replacements = {
        0: "FR5隧道喷涂运动与流体仿真项目\n合并论文",
        1: "运动仿真、路径规划、MoveIt2严格验证与独立流体喷雾仿真",
        3: "项目名称：隧道衬砌喷涂机器人运动仿真与流体喷雾仿真研究",
        4: "机器人平台：FAIRINO FR5六轴协作机器人",
        5: "研究内容：机器人算法运动仿真与流体喷雾仿真两条独立技术链",
        6: "日期：2026年7月",
    }
    for idx, text in replacements.items():
        replace_paragraph_text(doc.paragraphs[idx], text, size=14 if idx in (0, 1) else 12, bold=idx == 0)


def revise_final_delivery(doc):
    h = find_heading(doc, "1.10 第九版", "Heading 2")
    nxt = find_heading(doc, "1.11 MATLAB", "Heading 2")
    remove_between(h, nxt)
    replace_paragraph_text(h, "1.10 交付版（最终版）：开放拱形运动仿真与严格验收", size=14, bold=True)
    p = insert_after(h, "版本性质与交付边界：本节是机器人运动仿真的最终交付状态，不再称为普通“第九版”。交付对象包括开放拱形轨迹、181点连续逆运动学结果、MoveIt2规划输入、Ruckig平滑轨迹、正运动学复算、关节动力学和碰撞检查。流体仿真另列于第2章和第3章，使用独立OpenFOAM算例与后处理目录；两者可以共享经确认的几何或运动输入，但计算模型、状态变量、求解器和验收指标互不混用。")
    p = insert_after(p, "交付版运动仿真步骤如下。第一步，删除闭合马蹄形底部不参与喷涂的横线，形成开放拱形目标轮廓；第二步，依据墙面法向和150 mm虚拟TCP生成181个末端目标位姿；第三步，采用连续关节种子求解FR5逆运动学并检查关节限位；第四步，把轨迹送入ROS2与MoveIt2环境，完成规划、Ruckig时间参数化和FK复算；第五步，逐状态检查自碰撞、环境碰撞、相邻关节步长、速度、加速度和jerk；第六步，导出CSV、JSON、静态图和多视图动画；第七步，按“输入—求解—验证—输出”链条完成交付审查。")
    p = insert_after(p, "交付结果表明：输入四元数范数满足要求，法向最大对齐误差为8.54×10⁻⁷°，181个Ruckig平滑后状态未发现自碰撞或环境碰撞，开放轨迹的相邻关节转移均低于20°门限。早期闭合轨迹中存在的145.849°分支跳变，已通过13个喷涂段和12个停喷换姿段形成可审计工艺，其中喷涂段最大关节步长3.614°、换姿插值最大步长4.980°。这些结果证明当前算法链在虚拟TCP和简化环境下通过离线验收，但不等同于实体机器人已完成现场安全验证。")
    p = insert_after(p, "交付版与流体仿真的接口仅限于经过审查的TCP路径、速度和喷涂启停状态。流体模块读取这些信息时不得反向改变机器人关节求解、碰撞模型或MoveIt2验收结果；同理，运动仿真中的动画曲线、关节jerk和碰撞状态也不得解释为空气流场、颗粒扰动或沉积厚度。通过这种接口隔离，论文能够清楚回答“机器人如何运动”和“喷雾如何输运”两个不同问题，并在最终结论中分别陈述各自证据边界。")
    insert_after(p, "本交付版仍存在三项限制：其一，150 mm TCP是项目约定的虚拟设计参数，实体喷头安装后需重新标定；其二，碰撞环境没有包含软管、泵送设备和全部施工障碍物；其三，严格验收是离线软件结果，实机运行前仍需低速、停喷、限位和急停条件下的分级联调。因此，本版应表述为“运动仿真交付版”，而不是“机器人系统已投入生产”。")


def main():
    doc = Document(SRC)
    clean_cover(doc)
    revise_final_delivery(doc)
    add_version_workflows(doc)
    add_formula_explanations(doc)
    add_ai_chapter(doc)

    # Correct global wording that previously coupled the two models.
    for p in doc.paragraphs:
        if "运动学仿真" in p.text:
            for r in p.runs:
                if "运动学仿真" in r.text:
                    r.text = r.text.replace("运动学仿真", "运动仿真")
        p.paragraph_format.widow_control = True
        if p.style.name.startswith("Heading"):
            p.paragraph_format.keep_with_next = True
        if "最后一版粒子沉积热图" in p.text:
            for r in p.runs:
                r.text = r.text.replace("最后一版粒子沉积热图", "独立流体仿真沉积热图")
        if p.text.startswith("最后一版没有求解喷嘴内部流动"):
            replace_paragraph_text(
                p,
                "运动仿真交付版不求解喷嘴内部流动、空气相压力、湍流、颗粒相耦合或材料流变；这些内容由第2章和第3章的独立OpenFOAM流体仿真承担。两条技术链只通过经核验的几何、TCP路径、速度和喷涂启停状态形成单向数据接口，流体计算不得反向改变已经通过验收的关节轨迹和碰撞结论。这样的分离能够避免把机器人关节jerk误写成流体扰动，也避免把颗粒沉积分布误写成运动仿真结果。",
                size=12,
            )

    # The version summary must identify V9 as the motion-delivery state, not as
    # the separately documented fluid module.
    for table in doc.tables:
        for row in table.rows:
            if row.cells and row.cells[0].text.strip() == "V9" and len(row.cells) >= 4:
                values = [
                    "V9",
                    "运动仿真交付版（最终版）",
                    "开放拱形轨迹、连续IK、MoveIt2、Ruckig、FK与碰撞验证",
                    "形成181点开放轨迹及可审计的最终运动验收证据",
                ]
                for cell, value in zip(row.cells, values):
                    cell.text = value
                    for cp in cell.paragraphs:
                        for cr in cp.runs:
                            set_run_font(cr, size=9)

    normalize_captions(doc)
    strip_all_shading(doc)
    force_black_fonts(doc)

    # Compact the long automatic TOC so that its final page is not mostly blank.
    for style_name in ("toc 1", "toc 2", "toc 3"):
        if style_name in [s.name for s in doc.styles]:
            st = doc.styles[style_name]
            st.font.name = "Times New Roman"
            st._element.get_or_add_rPr().get_or_add_rFonts().set(qn("w:eastAsia"), "宋体")
            st.font.size = Pt(10)
            st.font.color.rgb = RGBColor(0, 0, 0)
            st.paragraph_format.space_before = Pt(0)
            st.paragraph_format.space_after = Pt(0)
            st.paragraph_format.line_spacing = 1.0

    # Body text uses black Song 12 pt / Times New Roman 12 pt; captions remain 9 pt.
    for p in doc.paragraphs:
        if re.match(r"^[图表]\d+-\d+$", p.text.strip()):
            continue
        for r in p.runs:
            if p.style.name == "Normal" and r.font.size is None:
                set_run_font(r, size=12)
    doc.core_properties.title = "FR5隧道喷涂运动与流体仿真论文规范修订版"
    doc.save(OUT)
    print(OUT)


if __name__ == "__main__":
    main()
