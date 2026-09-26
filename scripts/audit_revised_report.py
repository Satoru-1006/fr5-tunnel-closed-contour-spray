import re
from zipfile import ZipFile
from docx import Document

p = r"C:\Users\86198\Desktop\report_论文规范修订版.docx"
d = Document(p)
badcaps = [
    x.text for x in d.paragraphs
    if re.match(r"^[图表]\s*\d+", x.text.strip())
    and not re.match(r"^[图表]\d+-\d+$", x.text.strip())
]
flows = [x for x in d.paragraphs if x.text.startswith("本版性质界定")]
forms = [x for x in d.paragraphs if x.text.startswith("本节公式用于")]
v9 = [
    c.text for t in d.tables for row in t.rows
    if row.cells and row.cells[0].text.strip() == "V9" for c in row.cells
]
with ZipFile(p) as z:
    xml = z.read("word/document.xml").decode("utf-8")
colors = set(re.findall(r'<w:color[^>]*w:val="([^"]+)"', xml))
fills = set(re.findall(r'<w:shd[^>]*w:fill="([^"]+)"', xml))

print("paragraphs", len(d.paragraphs), "tables", len(d.tables), "images", len(d.inline_shapes))
print("version_boundaries", len(flows))
print("formula_explanations", len(forms), "min_formula_chars", min(map(lambda x: len(x.text), forms)))
print("bad_captions", len(badcaps), badcaps[:3])
print("V9", v9)
print("ai_chapter", sum(1 for x in d.paragraphs if x.text.startswith("第7章  AI辅助使用声明")))
print("old_particle_heading", sum(1 for x in d.paragraphs if x.style.name == "Heading 2" and "第九版" in x.text))
print("colors", colors)
print("fills", fills)
