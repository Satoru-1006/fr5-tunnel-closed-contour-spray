from __future__ import annotations

from pathlib import Path

from docx import Document
from docx.enum.section import WD_ORIENT
from docx.enum.table import WD_ALIGN_VERTICAL
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Cm, Pt, RGBColor

ROOT = Path(__file__).resolve().parents[1]
OUT = Path(r"C:\Users\86198\Desktop\word")
DOCX = OUT / "FR5开放拱形运动仿真_核心代码附录.docx"


def set_font(run, size=10.5, bold=False, color=None, code=False):
    name = "Consolas" if code else "Times New Roman"
    run.font.name = name
    run._element.rPr.rFonts.set(qn("w:ascii"), name)
    run._element.rPr.rFonts.set(qn("w:hAnsi"), name)
    run._element.rPr.rFonts.set(qn("w:eastAsia"), "宋体" if not code else "Consolas")
    run.font.size = Pt(size)
    run.bold = bold
    if color:
        run.font.color.rgb = RGBColor.from_string(color)


def add_heading(doc, text, size=14):
    p = doc.add_paragraph()
    p.alignment = WD_ALIGN_PARAGRAPH.LEFT
    p.paragraph_format.space_before = Pt(10)
    p.paragraph_format.space_after = Pt(5)
    p.paragraph_format.keep_with_next = True
    set_font(p.add_run(text), size=size, bold=True, color="17365D")


def add_body(doc, text, first=True):
    p = doc.add_paragraph()
    p.alignment = WD_ALIGN_PARAGRAPH.JUSTIFY
    p.paragraph_format.space_after = Pt(4)
    p.paragraph_format.line_spacing = 1.28
    if first:
        p.paragraph_format.first_line_indent = Cm(0.74)
    set_font(p.add_run(text), size=10.5)


def border(cell, top=None, bottom=None):
    tc_pr = cell._tc.get_or_add_tcPr()
    borders = tc_pr.first_child_found_in("w:tcBorders")
    if borders is None:
        borders = OxmlElement("w:tcBorders")
        tc_pr.append(borders)
    for edge, value in {"top": top, "bottom": bottom}.items():
        if value is None:
            continue
        node = borders.find(qn("w:" + edge))
        if node is None:
            node = OxmlElement("w:" + edge)
            borders.append(node)
        node.set(qn("w:val"), value)
        node.set(qn("w:sz"), "8")
        node.set(qn("w:color"), "000000")


def add_manifest(doc):
    p = doc.add_paragraph()
    p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    set_font(p.add_run("表A-1  代码附录文件清单"), size=9.5)
    rows = [
        ("离线轨迹", "scripts/generate_open_arch_lshape_animation.py", "构建181点开放拱形连续IK与基础动画"),
        ("多视角动画", "scripts/generate_open_arch_multiview.py", "输出左前、正前、右前真实六轴GIF"),
        ("MoveIt输入", "tools/build_open_arch_moveit_inputs.py", "世界坐标轨迹转换为base_link TCP姿态与种子关节"),
        ("输入校验", "ros2_moveit_bridge/validate_bridge_inputs.py", "校验开路径、四元数与法向数据"),
        ("严格链", "scripts/run_moveit_strict_validation.sh", "串联MoveIt、Ruckig、FK、动力学和碰撞验证"),
        ("规划器核心", "ros2_moveit_bridge/plan_closed_contour_moveit.py", "收录与开放路径相关的碰撞、定时及验证关键片段"),
        ("验收审计", "tools/audit_goal_requirements.py", "按目标条件生成最终审计结论"),
    ]
    table = doc.add_table(rows=1, cols=3)
    table.autofit = False
    widths = [Cm(2.6), Cm(8.3), Cm(12.0)]
    for i, text in enumerate(("模块", "源文件", "作用")):
        cell = table.rows[0].cells[i]
        cell.width = widths[i]; cell.vertical_alignment = WD_ALIGN_VERTICAL.CENTER; border(cell, top="single", bottom="single")
        cell.paragraphs[0].alignment = WD_ALIGN_PARAGRAPH.CENTER
        set_font(cell.paragraphs[0].add_run(text), size=8.8, bold=True)
    for row_i, values in enumerate(rows):
        cells = table.add_row().cells
        for i, text in enumerate(values):
            cells[i].width = widths[i]; cells[i].vertical_alignment = WD_ALIGN_VERTICAL.CENTER
            if row_i == len(rows) - 1:
                border(cells[i], bottom="single")
            p = cells[i].paragraphs[0]; p.paragraph_format.space_after = Pt(1)
            p.alignment = WD_ALIGN_PARAGRAPH.LEFT
            set_font(p.add_run(text), size=8.3, code=(i == 1))


def add_listing(doc, title, rel_path, start=None, end=None, note=None, page_break=True):
    path = ROOT / rel_path
    source = path.read_text(encoding="utf-8").splitlines()
    start = 1 if start is None else start
    end = len(source) if end is None else min(end, len(source))
    if page_break:
        doc.add_page_break()
    add_heading(doc, title, size=13)
    lead = f"文件：{rel_path}；收录范围：第 {start}–{end} 行。"
    if note:
        lead += note
    add_body(doc, lead, first=False)
    for line_no in range(start, end + 1):
        p = doc.add_paragraph()
        p.alignment = WD_ALIGN_PARAGRAPH.LEFT
        p.paragraph_format.left_indent = Cm(0.1)
        p.paragraph_format.first_line_indent = Cm(0)
        p.paragraph_format.space_after = Pt(0)
        p.paragraph_format.line_spacing = 1.0
        p.paragraph_format.keep_together = True
        line = f"{line_no:04d}  {source[line_no - 1]}"
        set_font(p.add_run(line), size=7.3, color="1F1F1F", code=True)


def build():
    OUT.mkdir(parents=True, exist_ok=True)
    doc = Document()
    section = doc.sections[0]
    section.orientation = WD_ORIENT.LANDSCAPE
    section.page_width = Cm(29.7); section.page_height = Cm(21.0)
    section.top_margin = Cm(1.7); section.bottom_margin = Cm(1.7); section.left_margin = Cm(1.7); section.right_margin = Cm(1.7)
    section.header_distance = Cm(0.9); section.footer_distance = Cm(0.9)
    header = section.header.paragraphs[0]; header.alignment = WD_ALIGN_PARAGRAPH.CENTER
    set_font(header.add_run("FR5 开放拱形喷涂运动仿真——核心代码附录"), size=8.5, color="666666")
    footer = section.footer.paragraphs[0]; footer.alignment = WD_ALIGN_PARAGRAPH.CENTER
    set_font(footer.add_run("附录A  核心源代码"), size=8.5, color="666666")

    p = doc.add_paragraph(); p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    p.paragraph_format.space_before = Pt(10); p.paragraph_format.space_after = Pt(8)
    set_font(p.add_run("附录A  FR5开放拱形喷涂运动仿真核心代码"), size=16, bold=True, color="17365D")
    p = doc.add_paragraph(); p.alignment = WD_ALIGN_PARAGRAPH.CENTER; p.paragraph_format.space_after = Pt(12)
    set_font(p.add_run("对应最终181点开放轨迹、真实六轴动画与MoveIt2严格验证链"), size=10.5, color="555555")
    add_body(doc, "本附录按最终版本的可复现流程组织源代码。完整收录离线连续逆运动学、三视角动画、MoveIt输入构建、输入校验、严格验证脚本与验收审计程序；MoveIt2规划器文件较长，故保留与本次开放路径改动直接相关的碰撞环境、弧长定时、动力学和碰撞验证片段。完整工程代码仍保留在项目源目录中。")
    add_manifest(doc)
    add_body(doc, "复现顺序为：先生成开放拱形连续IK及多视角动画，再构建base_link坐标系下的TCP姿态与关节种子，最后执行严格验证脚本。严格链将输出Ruckig后关节轨迹、FK TCP复算、动力学报告、碰撞报告和最终审计结论。", first=True)

    add_listing(doc, "A.1 开放拱形连续IK与单视角动画（完整文件）", "scripts/generate_open_arch_lshape_animation.py")
    add_listing(doc, "A.2 三种视角真实六轴GIF生成（完整文件）", "scripts/generate_open_arch_multiview.py")
    add_listing(doc, "A.3 开放拱形轨迹转换为MoveIt2输入（完整文件）", "tools/build_open_arch_moveit_inputs.py")
    add_listing(doc, "A.4 MoveIt2输入数据与开放路径校验（完整文件）", "ros2_moveit_bridge/validate_bridge_inputs.py")
    add_listing(doc, "A.5 MoveIt2严格验证执行脚本（完整文件）", "scripts/run_moveit_strict_validation.sh")
    add_listing(doc, "A.6 MoveIt2规划器：开放路径墙体碰撞环境（关键片段）", "ros2_moveit_bridge/plan_closed_contour_moveit.py", 664, 984, "该片段包含开放路径开关、TCP朝墙模式、墙体碰撞对象构建和弧长定时。")
    add_listing(doc, "A.7 MoveIt2规划器：动力学与碰撞验收（关键片段）", "ros2_moveit_bridge/plan_closed_contour_moveit.py", 1023, 1257, "该片段包含后Ruckig动力学门控与逐状态碰撞检查。")
    add_listing(doc, "A.8 目标验收审计程序（完整文件）", "tools/audit_goal_requirements.py", page_break=False)
    doc.core_properties.title = "FR5开放拱形运动仿真核心代码附录"
    doc.core_properties.subject = "最终开放轨迹与MoveIt2严格验证的核心源代码"
    doc.save(DOCX)
    print(DOCX)


if __name__ == "__main__":
    build()
