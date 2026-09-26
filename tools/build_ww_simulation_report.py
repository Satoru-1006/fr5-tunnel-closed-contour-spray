from __future__ import annotations

import json
import math
from pathlib import Path

from docx import Document
from docx.enum.section import WD_SECTION
from docx.enum.style import WD_STYLE_TYPE
from docx.enum.table import WD_CELL_VERTICAL_ALIGNMENT, WD_TABLE_ALIGNMENT
from docx.enum.text import WD_ALIGN_PARAGRAPH, WD_BREAK, WD_LINE_SPACING
from docx.oxml import OxmlElement, parse_xml
from docx.oxml.ns import nsdecls, qn
from docx.shared import Cm, Inches, Pt, RGBColor
from latex2mathml.converter import convert as latex_to_mathml
from lxml import etree


ROOT = Path(__file__).resolve().parents[1]
DATA_DIR = Path(r"C:\Users\86198\Desktop\ww")
WORK_DIR = ROOT / "tmp" / "ww_report"
FIG_DIR = WORK_DIR / "figures"
STATS_PATH = WORK_DIR / "derived_statistics.json"
OUT_PATH = DATA_DIR / "FR5带仰拱变截面隧道衬砌喷涂机器人运动仿真报告_图表紧凑排版修正版.docx"
MML2OMML = Path(r"C:\Program Files\Microsoft Office\root\Office16\MML2OMML.XSL")


CN_FONT = "宋体"
EN_FONT = "Times New Roman"
HEADING_FONT = "黑体"
INK = RGBColor(0x00, 0x00, 0x00)
MUTED = RGBColor(0x55, 0x55, 0x55)


def set_run_font(run, size=12, bold=False, italic=False, cn=CN_FONT, en=EN_FONT, color=INK):
    run.font.name = en
    run.font.size = Pt(size)
    run.font.bold = bold
    run.font.italic = italic
    run.font.color.rgb = color
    rpr = run._element.get_or_add_rPr()
    fonts = rpr.rFonts
    if fonts is None:
        fonts = OxmlElement("w:rFonts")
        rpr.insert(0, fonts)
    fonts.set(qn("w:ascii"), en)
    fonts.set(qn("w:hAnsi"), en)
    fonts.set(qn("w:eastAsia"), cn)
    fonts.set(qn("w:cs"), en)
    return run


def set_paragraph_format(paragraph, *, first_line=True, align=WD_ALIGN_PARAGRAPH.JUSTIFY, before=0, after=0, line=1.5):
    pf = paragraph.paragraph_format
    pf.alignment = align
    pf.space_before = Pt(before)
    pf.space_after = Pt(after)
    pf.line_spacing = line
    if first_line:
        pf.first_line_indent = Pt(24)
    else:
        pf.first_line_indent = Pt(0)


def add_body(doc, text, *, first_line=True, align=WD_ALIGN_PARAGRAPH.JUSTIFY, bold_prefix=None):
    p = doc.add_paragraph(style="正文")
    set_paragraph_format(p, first_line=first_line, align=align)
    if bold_prefix and text.startswith(bold_prefix):
        set_run_font(p.add_run(bold_prefix), bold=True)
        set_run_font(p.add_run(text[len(bold_prefix):]))
    else:
        set_run_font(p.add_run(text))
    return p


def add_heading(doc, text, level):
    p = doc.add_paragraph(style=f"Heading {level}")
    p.paragraph_format.keep_with_next = True
    p.paragraph_format.keep_together = True
    set_run_font(p.add_run(text), size={1: 18, 2: 15, 3: 14}[level], bold=True, cn=HEADING_FONT)
    return p


def set_repeat_table_header(row):
    tr_pr = row._tr.get_or_add_trPr()
    tbl_header = OxmlElement("w:tblHeader")
    tbl_header.set(qn("w:val"), "true")
    tr_pr.append(tbl_header)


def set_row_cant_split(row):
    """Prevent Word from splitting one table row across two pages."""
    tr_pr = row._tr.get_or_add_trPr()
    cant_split = tr_pr.find(qn("w:cantSplit"))
    if cant_split is None:
        tr_pr.append(OxmlElement("w:cantSplit"))


def set_cell_margins(cell, top=70, start=90, bottom=70, end=90):
    tc = cell._tc
    tc_pr = tc.get_or_add_tcPr()
    tc_mar = tc_pr.first_child_found_in("w:tcMar")
    if tc_mar is None:
        tc_mar = OxmlElement("w:tcMar")
        tc_pr.append(tc_mar)
    for m, v in (("top", top), ("start", start), ("bottom", bottom), ("end", end)):
        node = tc_mar.find(qn(f"w:{m}"))
        if node is None:
            node = OxmlElement(f"w:{m}")
            tc_mar.append(node)
        node.set(qn("w:w"), str(v))
        node.set(qn("w:type"), "dxa")


def set_cell_border(cell, **edges):
    tc_pr = cell._tc.get_or_add_tcPr()
    tc_borders = tc_pr.first_child_found_in("w:tcBorders")
    if tc_borders is None:
        tc_borders = OxmlElement("w:tcBorders")
        tc_pr.append(tc_borders)
    for edge in ("top", "left", "bottom", "right", "insideH", "insideV"):
        tag = f"w:{edge}"
        element = tc_borders.find(qn(tag))
        if element is None:
            element = OxmlElement(tag)
            tc_borders.append(element)
        cfg = edges.get(edge, {"val": "nil"})
        for key in ("val", "sz", "space", "color"):
            if key in cfg:
                element.set(qn(f"w:{key}"), str(cfg[key]))


def set_table_widths(table, widths_cm):
    table.autofit = False
    for row in table.rows:
        for cell, width in zip(row.cells, widths_cm):
            cell.width = Cm(width)
            tc_pr = cell._tc.get_or_add_tcPr()
            tc_w = tc_pr.first_child_found_in("w:tcW")
            if tc_w is None:
                tc_w = OxmlElement("w:tcW")
                tc_pr.append(tc_w)
            tc_w.set(qn("w:w"), str(int(width / 2.54 * 1440)))
            tc_w.set(qn("w:type"), "dxa")
    tbl_pr = table._tbl.tblPr
    tbl_w = tbl_pr.first_child_found_in("w:tblW")
    total = int(sum(widths_cm) / 2.54 * 1440)
    tbl_w.set(qn("w:w"), str(total))
    tbl_w.set(qn("w:type"), "dxa")


def add_three_line_table(doc, number, title, headers, rows, widths_cm, notes=None):
    cap = doc.add_paragraph(style="表题")
    cap.paragraph_format.keep_with_next = True
    cap.alignment = WD_ALIGN_PARAGRAPH.CENTER
    set_run_font(cap.add_run(f"表{number}  {title}"), size=10.5, bold=True)
    table = doc.add_table(rows=1, cols=len(headers))
    table.alignment = WD_TABLE_ALIGNMENT.CENTER
    table.autofit = False
    set_repeat_table_header(table.rows[0])
    for i, header in enumerate(headers):
        p = table.rows[0].cells[i].paragraphs[0]
        p.alignment = WD_ALIGN_PARAGRAPH.CENTER
        set_paragraph_format(p, first_line=False, align=WD_ALIGN_PARAGRAPH.CENTER, line=1.0)
        set_run_font(p.add_run(str(header)), size=10.5, bold=True)
    for row_values in rows:
        cells = table.add_row().cells
        for i, value in enumerate(row_values):
            p = cells[i].paragraphs[0]
            align = WD_ALIGN_PARAGRAPH.CENTER if i == 0 or len(str(value)) < 16 else WD_ALIGN_PARAGRAPH.LEFT
            set_paragraph_format(p, first_line=False, align=align, line=1.0)
            set_run_font(p.add_run(str(value)), size=10.5)
            cells[i].vertical_alignment = WD_CELL_VERTICAL_ALIGNMENT.CENTER
    set_table_widths(table, widths_cm)
    thin = {"val": "single", "sz": "6", "space": "0", "color": "000000"}
    thick = {"val": "single", "sz": "12", "space": "0", "color": "000000"}
    for r_idx, row in enumerate(table.rows):
        set_row_cant_split(row)
        for cell in row.cells:
            set_cell_margins(cell)
            # Keep every row linked to the next row.  The final row is linked
            # only when a table note follows.  Together with cantSplit this
            # makes a table that fits on one page move as one intact block.
            for paragraph in cell.paragraphs:
                paragraph.paragraph_format.keep_with_next = (
                    r_idx < len(table.rows) - 1 or bool(notes)
                )
            borders = {}
            if r_idx == 0:
                borders["top"] = thick
                borders["bottom"] = thin
            if r_idx == len(table.rows) - 1:
                borders["bottom"] = thick
            set_cell_border(cell, **borders)
    if notes:
        p = doc.add_paragraph(style="表注")
        set_paragraph_format(p, first_line=False, align=WD_ALIGN_PARAGRAPH.LEFT, line=1.0, before=2, after=4)
        set_run_font(p.add_run(f"注：{notes}"), size=9)
    return table


def add_figure(doc, filename, number, title, width_cm=15.6, note=None):
    p = doc.add_paragraph()
    p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    p.paragraph_format.keep_with_next = True
    p.paragraph_format.keep_together = True
    p.paragraph_format.space_before = Pt(4)
    p.paragraph_format.space_after = Pt(2)
    p.add_run().add_picture(str(FIG_DIR / filename), width=Cm(width_cm))
    cap = doc.add_paragraph(style="图题")
    cap.alignment = WD_ALIGN_PARAGRAPH.CENTER
    cap.paragraph_format.keep_with_next = bool(note)
    cap.paragraph_format.keep_together = True
    cap.paragraph_format.space_after = Pt(2 if note else 6)
    set_run_font(cap.add_run(f"图{number}  {title}"), size=10.5)
    if note:
        pn = doc.add_paragraph(style="图注")
        set_paragraph_format(pn, first_line=False, align=WD_ALIGN_PARAGRAPH.LEFT, line=1.0, after=6)
        pn.paragraph_format.keep_together = True
        set_run_font(pn.add_run(f"注：{note}"), size=9)
    return p


def latex_to_omml(latex: str):
    # The Office MML2OMML transform renders LaTeX bmatrix brackets as plain
    # one-line glyphs.  Express every bmatrix as an explicit scalable
    # delimiter around a matrix so Word creates m:d with full-height brackets.
    latex = latex.replace(r"\begin{bmatrix}", r"\left[\begin{matrix}")
    latex = latex.replace(r"\end{bmatrix}", r"\end{matrix}\right]")
    mathml = latex_to_mathml(latex)
    transform = etree.XSLT(etree.parse(str(MML2OMML)))
    omml = transform(etree.fromstring(mathml.encode("utf-8")))
    return parse_xml(etree.tostring(omml.getroot()))


def add_equation(doc, latex, number):
    p = doc.add_paragraph(style="公式")
    pf = p.paragraph_format
    pf.first_line_indent = Pt(0)
    pf.space_before = Pt(5)
    pf.space_after = Pt(5)
    pf.keep_together = True
    pf.tab_stops.add_tab_stop(Cm(8.0), 1)
    pf.tab_stops.add_tab_stop(Cm(16.0), 2)
    p.add_run("\t")
    p._p.append(latex_to_omml(latex))
    nr = p.add_run(f"\t（{number}）")
    set_run_font(nr, size=11)
    return p


def add_field(paragraph, instruction):
    run = paragraph.add_run()
    fld_char = OxmlElement("w:fldChar")
    fld_char.set(qn("w:fldCharType"), "begin")
    instr_text = OxmlElement("w:instrText")
    instr_text.set(qn("xml:space"), "preserve")
    instr_text.text = instruction
    fld_sep = OxmlElement("w:fldChar")
    fld_sep.set(qn("w:fldCharType"), "separate")
    text = OxmlElement("w:t")
    text.text = "更新域以显示内容"
    fld_end = OxmlElement("w:fldChar")
    fld_end.set(qn("w:fldCharType"), "end")
    run._r.extend([fld_char, instr_text, fld_sep, text, fld_end])


def add_page_number(paragraph):
    paragraph.alignment = WD_ALIGN_PARAGRAPH.CENTER
    r = paragraph.add_run("第 ")
    set_run_font(r, size=9)
    add_field(paragraph, "PAGE")
    set_run_font(paragraph.add_run(" 页"), size=9)


def configure_styles(doc):
    styles = doc.styles
    normal = styles["Normal"]
    normal.font.name = EN_FONT
    normal.font.size = Pt(12)
    normal._element.rPr.rFonts.set(qn("w:eastAsia"), CN_FONT)
    for name in ("正文", "表题", "表注", "图题", "图注", "公式", "参考文献"):
        if name not in styles:
            styles.add_style(name, WD_STYLE_TYPE.PARAGRAPH)
    body = styles["正文"]
    body.base_style = normal
    body.paragraph_format.first_line_indent = Pt(24)
    body.paragraph_format.line_spacing = 1.5
    body.paragraph_format.alignment = WD_ALIGN_PARAGRAPH.JUSTIFY
    body.paragraph_format.space_before = Pt(0)
    body.paragraph_format.space_after = Pt(0)
    for name in ("表题", "图题"):
        styles[name].base_style = normal
        styles[name].paragraph_format.alignment = WD_ALIGN_PARAGRAPH.CENTER
        styles[name].paragraph_format.line_spacing = 1.0
    for name in ("表注", "图注", "参考文献"):
        styles[name].base_style = normal
        styles[name].paragraph_format.line_spacing = 1.0
    styles["公式"].base_style = normal
    styles["公式"].paragraph_format.line_spacing = 1.0
    for level, size in ((1, 18), (2, 15), (3, 14)):
        style = styles[f"Heading {level}"]
        style.font.name = EN_FONT
        style.font.size = Pt(size)
        style.font.bold = True
        style.font.color.rgb = INK
        style._element.rPr.rFonts.set(qn("w:eastAsia"), HEADING_FONT)
        style.paragraph_format.alignment = WD_ALIGN_PARAGRAPH.LEFT
        style.paragraph_format.first_line_indent = Pt(0)
        style.paragraph_format.space_before = Pt(14 if level == 1 else 10)
        style.paragraph_format.space_after = Pt(8 if level == 1 else 5)
        style.paragraph_format.line_spacing = 1.0


def configure_document(doc):
    section = doc.sections[0]
    section.page_width = Cm(21.0)
    section.page_height = Cm(29.7)
    section.top_margin = Cm(2.54)
    section.bottom_margin = Cm(2.54)
    section.left_margin = Cm(3.0)
    section.right_margin = Cm(2.5)
    section.header_distance = Cm(1.5)
    section.footer_distance = Cm(1.5)
    section.different_first_page_header_footer = True
    header = section.header
    hp = header.paragraphs[0]
    hp.alignment = WD_ALIGN_PARAGRAPH.CENTER
    set_run_font(hp.add_run("FR5带仰拱变截面隧道衬砌喷涂机器人运动学仿真与路径规划设计"), size=9, color=MUTED)
    add_page_number(section.footer.paragraphs[0])


def add_cover(doc):
    for _ in range(3):
        doc.add_paragraph()
    p = doc.add_paragraph()
    p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    set_run_font(p.add_run("《机电综合实践》《机电系统课程设计》"), size=16, bold=True, cn=HEADING_FONT)
    p.paragraph_format.space_after = Pt(28)
    p = doc.add_paragraph()
    p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    set_run_font(p.add_run("课程设计论文报告"), size=24, bold=True, cn=HEADING_FONT)
    p.paragraph_format.space_after = Pt(36)
    p = doc.add_paragraph()
    p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    set_run_font(p.add_run("FR5带仰拱变截面隧道衬砌\n喷涂机器人运动仿真\n建立方法与结果分析"), size=22, bold=True, cn=HEADING_FONT)
    p.paragraph_format.line_spacing = 1.4
    p.paragraph_format.space_after = Pt(56)
    metadata = [
        ("项目名称", "带仰拱变截面隧道衬砌喷涂机器人运动学仿真与路径规划设计"),
        ("机器人平台", "法奥 FR5 六轴协作机器人"),
        ("学生", "侯冰洋（组长，23019406）；赵晓钧（23019423）"),
        ("指导教师", "丁彦玉"),
        ("报告日期", "2026年7月"),
    ]
    table = doc.add_table(rows=len(metadata), cols=2)
    table.alignment = WD_TABLE_ALIGNMENT.CENTER
    set_table_widths(table, [3.1, 10.5])
    for row, (label, value) in zip(table.rows, metadata):
        for cell in row.cells:
            set_cell_border(cell)
            set_cell_margins(cell, top=100, bottom=100)
        p0 = row.cells[0].paragraphs[0]
        p0.alignment = WD_ALIGN_PARAGRAPH.RIGHT
        set_run_font(p0.add_run(label + "："), size=13, bold=True, cn=HEADING_FONT)
        p1 = row.cells[1].paragraphs[0]
        p1.alignment = WD_ALIGN_PARAGRAPH.LEFT
        set_run_font(p1.add_run(value), size=13)
    doc.add_page_break()


def add_abstracts(doc):
    add_heading(doc, "摘  要", 1)
    add_body(
        doc,
        "针对带曲线仰拱变截面马蹄形隧道衬砌喷涂作业中截面尺寸沿轴向变化、喷枪姿态连续性和机器人避碰约束相互耦合的问题，本文依据项目文件夹中的源程序、URDF 模型、逆运动学路径点、时间参数化轨迹、验证记录和多视角动画，对 FR5 带仰拱变截面隧道喷涂运动仿真的建立过程进行了可复核重构。仿真首先以纵向位置为自变量建立宽度、高度及仰拱深度周期变化的封闭马蹄形截面，并在内壁法向上设置 50 mm 喷涂间距；随后引入 150 mm 虚拟 TCP，将喷涂位姿换算为法兰目标位姿。机器人模型读取官方 FR5 URDF 和 MoveIt2 关节约束，采用带边界的非线性最小二乘法求解逆运动学，并通过多初值搜索、关节连续性代价和离散连杆净空检查筛选候选解。最后采用分段 Ruckig 算法进行速度、加速度与加加速度约束下的时间参数化，输出 9900 个时序采样点及多视角动画。",
    )
    add_body(
        doc,
        "结果表明，仿真链条能够生成 5 个纵向站位、285 个逆运动学路径点和 390.59 s 的可视化运动序列；最大逆运动学位置误差为 3.00 mm，时间参数化后最大相邻关节采样步长为 4.09°。然而，实际 TCP 相对规划插值路径的最大偏差达到 510.03 mm，1043 个采样点超过 10 mm，离散连杆净空最小值为 −197.61 mm，并记录到 4 个碰撞采样点。因此，该文件夹证明了运动仿真框架和数据输出链已经建立，但当前结果仍处于“review_required”状态，不能直接作为实体喷涂执行轨迹。本文据此提出统一喷涂路径与机器人实际 TCP 语义、消除逆解分支跃迁、在时间参数化阶段加入笛卡尔跟踪约束以及使用完整碰撞几何复核等改进方向。",
    )
    p = doc.add_paragraph(style="正文")
    set_paragraph_format(p, first_line=False, align=WD_ALIGN_PARAGRAPH.LEFT)
    set_run_font(p.add_run("关键词："), bold=True)
    set_run_font(p.add_run("FR5 协作机器人；带仰拱变截面隧道；隧道衬砌喷涂；逆运动学；轨迹规划；Ruckig；碰撞检测"))

    add_heading(doc, "ABSTRACT", 1)
    p = doc.add_paragraph(style="正文")
    set_paragraph_format(p, first_line=True)
    set_run_font(
        p.add_run(
            "This report reconstructs an auditable FR5 tunnel-lining spraying simulation for a longitudinally variable horseshoe tunnel with a curved invert, using the source code, URDF model, inverse-kinematics waypoints, time-parameterized trajectory, validation record and multi-view animations contained in the project folder. The variable cross-section is parameterized through its width, height and invert depth. TCP targets are offset by 50 mm from the wall and converted into flange poses using a 150 mm virtual tool. The official FR5 URDF and MoveIt2 joint limits are then used in bounded nonlinear least-squares inverse kinematics with deterministic multi-seed search, continuity scoring and sampled link-clearance gating. Segmented Ruckig retiming finally produces 9,900 samples over 390.59 s. The maximum IK position error is 3.00 mm and the maximum adjacent time-sample joint step is 4.09 degrees. Nevertheless, the maximum realized TCP path error is 510.03 mm, 1,043 samples exceed 10 mm, and four sampled collision states are reported. The simulation pipeline is therefore functionally established but not yet an executable, collision-free production trajectory."
        ),
        size=12,
        cn=EN_FONT,
    )
    p = doc.add_paragraph(style="正文")
    set_paragraph_format(p, first_line=False, align=WD_ALIGN_PARAGRAPH.LEFT)
    set_run_font(p.add_run("Key words: "), bold=True, cn=EN_FONT)
    set_run_font(p.add_run("FR5 collaborative robot; variable cross-section tunnel with invert; tunnel lining spraying; inverse kinematics; trajectory planning; Ruckig; collision checking"), cn=EN_FONT)
    doc.add_page_break()


def add_toc(doc):
    p = doc.add_paragraph()
    p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    set_run_font(p.add_run("目  录"), size=18, bold=True, cn=HEADING_FONT)
    p.paragraph_format.space_after = Pt(12)
    entries = [
        ("第1章  绪论", "5"),
        ("第2章  带仰拱变截面隧道与喷涂路径几何建模", "6"),
        ("第3章  FR5机器人运动学模型与TCP定义", "10"),
        ("第4章  逆运动学、分支选择与碰撞门禁", "12"),
        ("第5章  Ruckig时间参数化与动画建立", "15"),
        ("第6章  仿真结果与定量分析", "19"),
        ("第7章  讨论与改进建议", "24"),
        ("第8章  结论", "26"),
        ("参考文献", "27"),
        ("附录A  仿真复现与输出说明", "27"),
        ("附录B  人工智能辅助使用说明", "28"),
    ]
    for title, page in entries:
        p = doc.add_paragraph()
        set_paragraph_format(p, first_line=False, align=WD_ALIGN_PARAGRAPH.LEFT, line=1.3, after=2)
        p.paragraph_format.tab_stops.add_tab_stop(Cm(16.0), 2, 1)
        set_run_font(p.add_run(f"{title}\t{page}"), size=12)
    doc.add_page_break()


def build_report():
    stats = json.loads(STATS_PATH.read_text(encoding="utf-8"))
    v = stats["validation"]
    doc = Document()
    configure_styles(doc)
    configure_document(doc)
    add_cover(doc)
    add_abstracts(doc)
    add_toc(doc)

    add_heading(doc, "第1章  绪论", 1)
    add_heading(doc, "1.1 项目背景与任务边界", 2)
    add_body(doc, "隧道衬砌喷涂机器人需要在狭窄、粉尘和高湿环境中持续保持喷枪与壁面之间的距离和夹角，同时避免机械臂本体与衬砌发生干涉。相较于普通点到点搬运，喷涂任务不仅要求末端到达离散位置，还要求沿连续曲面建立位置、姿态、速度和喷涂状态之间的一致映射。因此，运动学建模、全覆盖路径离散、逆运动学分支选择、关节轨迹平滑和碰撞检查必须作为同一条数据链处理。")
    add_body(doc, "任务书要求以法奥 FR5 六轴协作机器人为核心，完成正逆运动学可视化、至少两类隧道截面的喷涂轨迹调试、虚拟样机与控制环境搭建，并进一步覆盖末端工具、流体仿真、机械电气设计和上位机等内容。本文仅对当前文件夹中已有的带曲线仰拱变截面运动仿真成果进行论文级重构和评价，不把文件夹外尚无证据的流体仿真、结构图、电气图或 Qt 上位机功能写成已完成成果。")
    add_heading(doc, "1.2 本文研究内容", 2)
    add_body(doc, "本文围绕四个问题展开：第一，变截面马蹄形隧道如何由少量参数建立；第二，壁面法向、喷涂间距与虚拟 TCP 如何转化为 FR5 法兰目标位姿；第三，数值逆运动学、碰撞门禁和分段 Ruckig 时间参数化如何串联；第四，已有 CSV、JSON 和 GIF 结果究竟证明了什么，又暴露了哪些工程风险。图1-1给出了从几何输入到可审计输出的完整计算流程。")
    add_figure(doc, "fig_1_1_workflow.png", "1-1", "FR5 变截面隧道喷涂运动仿真数据链", 15.8, "流程图依据项目源程序的调用关系绘制，非生成式图片。")
    add_body(doc, "由图1-1可见，仿真不是单独播放一段机器人动画，而是由隧道几何、TCP目标、FR5运动学、逆解与碰撞筛选、Ruckig重定时以及结果输出六个环节顺序构成。前一环节的输出直接成为后一环节的输入，因此任何几何定义、坐标变换或逆解分支错误都会继续传递至最终CSV、JSON和动画。该流程图的作用是明确后续各章之间的逻辑关系，而不是作为脱离正文的装饰性图片。")
    add_heading(doc, "1.3 数据来源与可追溯性", 2)
    add_body(doc, "为避免只依据动画外观评价仿真质量，需要先说明论文采用了哪些原始文件，以及各文件能够支持哪一类结论。表1.1按照“文件—数据规模或类型—证据角色”三个层次汇总当前文件夹中的主要输入和输出，使后续公式计算、图形重绘与验证结论均能够追溯到具体数据源。")
    add_three_line_table(
        doc,
        "1.1",
        "论文使用的项目文件与证据角色",
        ["文件", "数据规模/类型", "证据角色"],
        [
            ["variable_invert_ik_waypoints.csv", "285×7", "逆运动学路径点及六关节角"],
            ["variable_invert_tcp_poses.csv", "9900×11", "时间、实际 TCP 坐标及六关节角"],
            ["variable_invert_validation.json", "1 个验证记录", "参数、误差、碰撞和状态结论"],
            ["variable_invert_geometry.png", "截面图", "原始几何可视化成果"],
            ["variable_invert_motion*.gif", "6 个动画", "单视角、三视角和四视角运动证据"],
            ["run_fr5_invert_variable_section.py", "主程序", "几何、IK、碰撞、重定时和输出逻辑"],
            ["fairino5_v6.urdf", "官方机器人模型", "六关节坐标链和关节边界"],
        ],
        [5.8, 3.1, 7.1],
        "所有统计图均由上述 CSV/JSON 数据或同一源程序中的解析几何方程计算得到。",
    )
    add_body(doc, "由表1.1可知，variable_invert_ik_waypoints.csv用于描述时间参数化前的285个逆运动学路径点，variable_invert_tcp_poses.csv则保存9900个时序采样点，是关节曲线、TCP轨迹和速度分析的直接数据来源。validation.json集中记录误差、碰撞和综合状态，URDF与主程序分别提供机器人几何链和计算逻辑。因而，本报告中的图形和结论不是从单一图片推测得到，而是由源程序、数值数据和验证记录相互交叉支撑。")

    add_heading(doc, "第2章  带仰拱变截面隧道与喷涂路径几何建模", 1)
    add_heading(doc, "2.1 纵向变截面参数化", 2)
    add_body(doc, "仿真长度取 L=0.45 m，并沿隧道轴向设置 5 个站位。为使截面尺寸随纵向位置平滑变化，程序采用正弦与余弦函数分别描述宽度、高度和仰拱深度。设 w0 为基准宽度、Aw 为宽度变化幅值，纵向坐标为 y，则宽度的原始参数化关系式为")
    add_equation(doc, r"w(y)=w_0+A_w\sin\left(2\pi\frac{y}{L}\right)", "2.1")
    add_body(doc, "本算例取 w0=0.78 m、Aw=0.08 m。以第二站位 y=0.1125 m 为例，由式（2.1）可得")
    add_equation(doc, r"w(0.1125)=0.78+0.08\sin\left(2\pi\frac{0.1125}{0.45}\right)=0.860\ \mathrm m", "2.1a")
    add_body(doc, "计算结果与表2.1中第二站位的宽度 0.860 m 一致。由于正弦项的取值范围为[−1,1]，式（2.1）还给出宽度理论范围 0.78−0.08=0.70 m 至 0.78+0.08=0.86 m；入口和出口处正弦项均为零，因此两端宽度均回到 0.78 m。对应地，设 h0 为基准高度、Ah 为高度变化幅值，则截面高度的原始关系式为")
    add_equation(doc, r"h(y)=h_0+A_h\cos\left(2\pi\frac{y}{L}\right)", "2.2")
    add_body(doc, "本算例取 h0=0.76 m、Ah=0.06 m。中间站位 y=0.225 m 对应余弦函数的最低点，由式（2.2）可得")
    add_equation(doc, r"h(0.225)=0.76+0.06\cos\left(2\pi\frac{0.225}{0.45}\right)=0.700\ \mathrm m", "2.2a")
    add_body(doc, "式（2.2a）表明中间站位的净高为 0.700 m，而入口和出口处高度均为 0.820 m。宽度和高度采用不同的三角函数相位，使截面不只是整体等比例缩放，而是在纵向形成宽、高相互耦合的变截面。设 d0 为基准仰拱深度、Ad 为附加变化幅值，则曲线仰拱深度的原始关系式为")
    add_equation(doc, r"d_{\mathrm{inv}}(y)=d_0+A_d\sin\left(\pi\frac{y}{L}\right)", "2.3")
    add_body(doc, "本算例取 d0=0.030 m、Ad=0.015 m。仍以中间站位 y=0.225 m 为例，由式（2.3）可得")
    add_equation(doc, r"d_{\mathrm{inv}}(0.225)=0.030+0.015\sin\left(\pi\frac{0.225}{0.45}\right)=0.045\ \mathrm m", "2.3a")
    add_body(doc, "由式（2.3a）可知，仰拱底部在中间站位下沉 45 mm，两端则为 30 mm。宽度、高度和仰拱深度三项参数均连续可导，因此相邻纵向站位之间不存在由参数突变引起的几何折断。三项参数的连续变化如图2-1所示，离散站位在曲线上以圆点标出。")
    add_figure(doc, "fig_2_2_section_parameters.png", "2-1", "截面宽度、高度和仰拱深度沿隧道轴向的变化", 15.8)
    add_body(doc, "图2-1表明，截面宽度在 y=0.1125 m 附近达到0.860 m、在 y=0.3375 m 附近降低至0.700 m；截面高度则在中部降至0.700 m，并在两端恢复至0.820 m。仰拱深度从两端的0.030 m平滑增加至中部的0.045 m。三条曲线均连续变化，说明当前模型属于宽度、高度和仰拱深度共同变化的纵向变截面，而不是将同一个固定截面简单平移五次。")
    add_heading(doc, "2.2 马蹄形截面构造", 2)
    add_body(doc, "单个截面由左直墙、上半圆拱、右直墙和下部椭圆仰拱按固定顺序拼接。设 r=w/2，拱脚高度 s=h−r，参数角 α 从 0 变化至 π，则拱顶曲线的原始参数方程为")
    add_equation(doc, r"x=r\cos\alpha,\quad z=s+r\sin\alpha,\quad \alpha\in[0,\pi]", "2.4")
    add_body(doc, "在中间站位 y=0.225 m 处，w=0.780 m、h=0.700 m，因而 r=0.390 m、s=0.700−0.390=0.310 m。取拱顶参数 α=π/2，由式（2.4）可得")
    add_equation(doc, r"x=0.390\cos\frac{\pi}{2}=0,\quad z=0.310+0.390\sin\frac{\pi}{2}=0.700\ \mathrm m", "2.4a")
    add_body(doc, "式（2.4a）重新得到该站位 0.700 m 的拱顶高度，说明由宽度、高度派生的圆拱半径和拱脚高度在几何上闭合。下部仰拱采用在左右拱脚处具有竖直切向的半椭圆表达式")
    add_equation(doc, r"x=r\cos\beta,\quad z=-d_{\mathrm{inv}}\sin\beta,\quad \beta\in[0,\pi]", "2.5")
    add_body(doc, "中间站位的仰拱深度为 0.045 m。取 β=π/2 对应仰拱底部中点，由式（2.5）可得")
    add_equation(doc, r"x=0.390\cos\frac{\pi}{2}=0,\quad z=-0.045\sin\frac{\pi}{2}=-0.045\ \mathrm m", "2.5a")
    add_body(doc, "式（2.5a）说明该站位仰拱底点位于拱脚基准面以下 45 mm。根据式（2.4）和式（2.5），上拱、直墙与仰拱可形成闭合轮廓。程序将轮廓起点滚动到仰拱底部中点，运动顺序依次为底部中点、左仰拱、左直墙、拱顶、右直墙、右仰拱并返回底部中点。这一处理使每个截面的喷涂过程具有明确起止语义。五个站位的截面和 TCP 内偏轮廓见图2-2。")
    add_figure(doc, "fig_2_1_variable_sections.png", "2-2", "五个纵向站位的隧道内壁与 TCP 喷涂轮廓", 12.3, "灰线为隧道截面；红线为根据 50 mm 离壁距离构造的显示轮廓。图例位于数据区外。")
    add_body(doc, "由图2-2可见，五个站位均保持“直墙—圆拱—曲线仰拱”的闭合马蹄形拓扑，但轮廓宽度、拱顶高度和仰拱下沉量随纵向位置发生变化。红色TCP轮廓整体位于灰色隧道内壁的内侧，直观体现了50 mm离壁要求。为了把图中的几何趋势转化为可复核数值，表2.1进一步列出五个站位的具体参数。")
    add_three_line_table(
        doc,
        "2.1",
        "五个离散站位的截面参数",
        ["站位 y/m", "宽度 w/m", "高度 h/m", "仰拱深度 d/m"],
        [["0.000", "0.780", "0.820", "0.0300"], ["0.113", "0.860", "0.760", "0.0406"], ["0.225", "0.780", "0.700", "0.0450"], ["0.338", "0.700", "0.760", "0.0406"], ["0.450", "0.780", "0.820", "0.0300"]],
        [4.0, 4.0, 4.0, 4.0],
    )
    add_body(doc, "表2.1与图2-1、图2-2相互对应：第二站位宽度最大，为0.860 m；第四站位宽度最小，为0.700 m；第三站位高度最低且仰拱最深，分别为0.700 m和0.0450 m。入口与出口的三项参数完全相同，说明所选三角函数在0.45 m建模区间两端实现了尺寸闭合。该表同时为后续壁面法向计算和碰撞边界计算提供逐站位参数依据。")
    add_heading(doc, "2.3 壁面法向与喷涂间距", 2)
    add_body(doc, "对闭合轮廓相邻点求数值梯度获得切向量 t，经归一化后构造二维候选法向，并通过截面内部参考点判断符号。设单位内法向为 n，则目标 TCP 位置由壁面点 p_w 沿 n 偏置 50 mm 得到：")
    add_equation(doc, r"\boldsymbol p_{\mathrm{tcp}}=\boldsymbol p_w+d_s\boldsymbol n,\qquad d_s=0.05\ \mathrm m", "2.6")
    add_body(doc, "以中间站位拱顶为例，计入程序施加的 0.14 m 竖向平移后，壁面点为 p_w=[0,0.225,0.840]ᵀ m，指向截面内部的单位法向为 n=[0,0,−1]ᵀ。由式（2.6）可得")
    add_equation(doc, r"\boldsymbol p_{\mathrm{tcp}}=\begin{bmatrix}0\\0.225\\0.840\end{bmatrix}+0.05\begin{bmatrix}0\\0\\-1\end{bmatrix}=\begin{bmatrix}0\\0.225\\0.790\end{bmatrix}\ \mathrm m", "2.6a")
    add_body(doc, "式（2.6a）表明 TCP 相对拱顶壁面沿内法向回撤 50 mm，欧氏距离正好为 0.05 m。该计算给出了喷涂间距的直接几何意义。需要指出，程序用于动画显示的红色轮廓是解析缩小截面，而逆运动学目标使用数值法向偏置，两者在拐接区域并非严格同一个离散几何对象。这种“显示路径”与“求解路径”的语义差异，是后续实际 TCP 偏差分析必须关注的问题。")

    add_heading(doc, "第3章  FR5 机器人运动学模型与 TCP 定义", 1)
    add_heading(doc, "3.1 官方 URDF 模型加载", 2)
    add_body(doc, "项目没有使用临时 DH 占位模型，而是从 fairino5_v6.urdf 读取六个转动关节的原点、姿态、转轴和位置边界，并从 MoveIt2 的 joint_limits.yaml 读取最大速度和最大加速度。FR5 官方产品资料给出的基本指标包括 6 个转动自由度、额定负载 5 kg、最大工作半径 922 mm、重复定位精度 ±0.02 mm。为区分产品规格、仿真布置参数和当前建模假设，表3.1对这些数值及其来源进行集中说明。")
    add_three_line_table(
        doc,
        "3.1",
        "FR5 平台与本仿真关键参数",
        ["参数", "数值", "来源/用途"],
        [["自由度", "6", "FR5 官方产品资料及 URDF"], ["额定负载", "5 kg", "平台规格说明"], ["最大工作半径", "922 mm", "平台规格说明"], ["重复定位精度", "±0.02 mm", "ISO 9283 平台指标"], ["机器人基座位置", "(0, −0.25, 0.20) m", "仿真源程序"], ["基座偏航角", "π rad", "面向隧道内部"], ["虚拟 TCP 长度", "150 mm", "验证记录；当前为假定值"], ["喷涂离壁距离", "50 mm", "路径规划参数"]],
        [4.2, 4.0, 7.8],
        "150 mm 虚拟 TCP 是本仿真的建模约定，不代表已经完成真实喷枪标定。",
    )
    add_body(doc, "由表3.1可知，机器人自由度、额定负载和工作半径属于平台层面的基本能力，而基座位置、基座偏航角、虚拟TCP长度和喷涂离壁距离属于本次仿真的场景参数。尤其是150 mm虚拟TCP只用于构造当前法兰目标，不能与真实喷枪标定结果等同。后续正运动学和逆运动学计算均以表3.1中的基座布置和工具长度为前提，因此这些参数发生变化时必须重新计算全部轨迹。")
    add_heading(doc, "3.2 正运动学", 2)
    add_body(doc, "对于第 i 个转动关节，URDF 给出固定原点变换 T_origin,i 和绕关节轴旋转的变量变换 T_axis,i(q_i)。机器人法兰相对世界坐标系的齐次变换为")
    add_equation(doc, r"{}^{W}\boldsymbol T_F(\boldsymbol q)={}^W\boldsymbol T_B\prod_{i=1}^{6}\left(\boldsymbol T_{\mathrm{origin},i}\boldsymbol T_{\mathrm{axis},i}(q_i)\right)", "3.1")
    add_body(doc, "式（3.1）是串联六个关节固定变换和转动变换的原始正运动学公式。以 variable_invert_ik_waypoints.csv 的第 1 个关节状态为例，实际关节角代入为")
    add_equation(doc, r"\boldsymbol q_0=\begin{bmatrix}1.144063&-0.944300&2.308499&0.209736&1.563895&0.242601\end{bmatrix}^{\mathrm T}\ \mathrm{rad}", "3.1a")
    add_body(doc, "将式（3.1a）的六个角度依次代入 URDF 关节链，并使用基座位置 [0,−0.25,0.20]ᵀ m 及绕 z 轴 π rad 的基座偏航角，由式（3.1）计算得到法兰齐次变换")
    add_equation(doc, r"{}^{W}\boldsymbol T_F(\boldsymbol q_0)=\begin{bmatrix}-0.784208&0.620451&-0.007577&0.001513\\0.620469&0.784231&0&0.000004\\0.005942&-0.004701&-0.999971&0.309997\\0&0&0&1\end{bmatrix}", "3.1b")
    add_body(doc, "式（3.1b）的最后一列给出法兰位置 [0.001513,0.000004,0.309997]ᵀ m，左上 3×3 子矩阵给出法兰姿态。该数值由当前官方 URDF 和实际 CSV 首行关节角复算得到，而不是示意值。式（3.1）在程序的 fk() 与 all_joint_frames() 中实现：前者求取法兰位姿，后者保存基座到各关节的全部坐标帧，用于动画连杆绘制和离散碰撞检查。")
    add_heading(doc, "3.3 虚拟 TCP 与喷枪姿态", 2)
    add_body(doc, "工具坐标系相对法兰沿局部 z 轴平移 0.15 m，因此 TCP 变换为")
    add_equation(doc, r"{}^{W}\boldsymbol T_{\mathrm{TCP}}={}^W\boldsymbol T_F\,{}^F\boldsymbol T_{\mathrm{TCP}},\qquad {}^F\boldsymbol p_{\mathrm{TCP}}=[0,0,0.15]^\mathrm T", "3.2")
    add_body(doc, "由式（3.2）可知，TCP 位置等于法兰位置加上法兰局部 z 轴方向的 0.15 m 平移。式（3.1b）第三列为法兰局部 z 轴在世界坐标系中的方向，因此首个采样点的实际代入计算为")
    add_equation(doc, r"{}^W\boldsymbol p_{\mathrm{TCP}}=\begin{bmatrix}0.001513\\0.000004\\0.309997\end{bmatrix}+0.15\begin{bmatrix}-0.007577\\0\\-0.999971\end{bmatrix}=\begin{bmatrix}0.000376\\0.000004\\0.160001\end{bmatrix}\ \mathrm m", "3.2a")
    add_body(doc, "式（3.2a）的结果与 variable_invert_tcp_poses.csv 首行记录 [0.000376,0.000004,0.160001]ᵀ m 一致，从而验证了“关节角—法兰位姿—虚拟 TCP”计算链在该采样点上的数值闭合性。150 mm 仅为当前建模约定，真实喷枪安装后仍需通过工具坐标标定替换。")
    add_body(doc, "为使喷枪轴线朝向衬砌壁面，程序令目标坐标系的局部 −z 轴沿壁面法向反向，并用隧道轴向单位向量作为辅助方向。其正交基构造为")
    add_equation(doc, r"\boldsymbol z=-\boldsymbol n,\quad \boldsymbol x=\frac{\boldsymbol e_y\times\boldsymbol z}{\left\|\boldsymbol e_y\times\boldsymbol z\right\|},\quad \boldsymbol y=\boldsymbol z\times\boldsymbol x", "3.3")
    add_body(doc, "以仰拱底部中点为例，其指向隧道内部的单位法向为 n=[0,0,1]ᵀ，隧道轴向辅助向量为 e_y=[0,1,0]ᵀ。由式（3.3）可得目标工具坐标基")
    add_equation(doc, r"\boldsymbol z=\begin{bmatrix}0\\0\\-1\end{bmatrix},\quad \boldsymbol x=\begin{bmatrix}-1\\0\\0\end{bmatrix},\quad \boldsymbol y=\begin{bmatrix}0\\1\\0\end{bmatrix}", "3.3a")
    add_body(doc, "式（3.3a）表示喷枪局部 z 轴朝向仰拱壁面，而局部 y 轴保持与隧道纵向一致。对其他轮廓点重复该叉乘和归一化过程，TCP 姿态即可随截面法向变化。由于逆运动学求解对象是法兰位姿，程序通过目标 TCP 变换右乘工具变换的逆矩阵得到法兰目标，即 T_F,target=T_TCP,target T_tool^−1。")

    add_heading(doc, "第4章  逆运动学、分支选择与碰撞门禁", 1)
    add_heading(doc, "4.1 带边界非线性最小二乘逆解", 2)
    add_body(doc, "逆运动学采用 SciPy 的带边界非线性最小二乘求解器。残差同时包含法兰位置误差、工具 z 轴方向误差、可选完整姿态误差和相对初值的连续性代价。其目标可写为")
    add_equation(doc, r"\min_{\boldsymbol q_{\min}\le\boldsymbol q\le\boldsymbol q_{\max}}\left\|\begin{bmatrix}\boldsymbol p(\boldsymbol q)-\boldsymbol p_d\\ w_o(\boldsymbol z(\boldsymbol q)-\boldsymbol z_d)\\ w_R\boldsymbol e_R\\ w_c(\boldsymbol q-\boldsymbol q_s)\end{bmatrix}\right\|_2^2", "4.1")
    add_body(doc, "式（4.1）中，本算例对方向轴误差取 w_o=1.0，对连续性代价在常规候选中取 w_c=0.001，完整姿态旋转向量默认不参与残差，即 w_R=0。以首个路径点为例，正运动学复算得到法兰位置误差范数 0.001513 m、工具轴方向误差范数 0.007577，候选解相对初始种子的关节差范数为 1.6273 rad。将这些真实数据代入式（4.1）可得近似目标函数值")
    add_equation(doc, r"J_0=(0.001513)^2+(1.0\times0.007577)^2+(0.001\times1.6273)^2=6.23\times10^{-5}", "4.1a")
    add_body(doc, "式（4.1a）说明该点残差主要由工具轴方向项贡献，而连续性项因权重仅为 0.001，对单点目标函数的贡献较小。这也解释了为什么局部逆解能够获得较小位置误差，却仍可能在不同路径点之间切换关节分支。在额外搜索和回退阶段，程序把 w_c 提高至 0.004，但最终仍需通过相邻关节步长和碰撞指标独立审查。")
    add_heading(doc, "4.2 多初值搜索与分支筛选", 2)
    add_body(doc, "六轴机器人对同一末端位姿通常存在多个肘部和腕部构型。程序对每个目标点依次尝试固定种子、上一点解、备用构型、四组确定性扰动和八组带固定随机种子的扰动；若局部候选均与隧道发生干涉，再增加 32 组更宽范围的初值。对候选关节差采用 2π 周期内的最短等效差")
    add_equation(doc, r"\Delta q_j=\operatorname{wrap}_{[-\pi,\pi)}\left(q_{j,k}-q_{j,k-1}\right)", "4.2")
    add_body(doc, "根据式（4.2），程序优先保留最大关节变化不超过 25° 的候选，并在候选误差与连续性范数之间进行加权选择。以时间参数化后的最大相邻变化为例，CSV 第 7548 和 7549 个采样点的第5关节角分别为 0.408492 rad 和 0.337169 rad，代入式（4.2）得到")
    add_equation(doc, r"\Delta q_5=\operatorname{wrap}(0.337169-0.408492)=-0.071322\ \mathrm{rad}=-4.086^\circ", "4.2a")
    add_body(doc, "式（4.2a）的绝对值 4.086° 与 validation.json 中 max_adjacent_joint_step_deg=4.086468° 一致，并低于时间采样轨迹采用的 20°判据。图4-1给出了 285 个原始 IK 路径点的六关节角。需要强调，时间采样后的局部步长通过，并不能消除原始 IK 路径点之间的分支跃迁。")
    add_body(doc, "为便于识别分支切换，图4-1采用相同的横坐标索引分别展示六个关节，并在各子图中叠加URDF位置边界。阅读时首先比较曲线是否接近关节上下限，其次检查相邻路径点是否出现近似竖直的跳变，最后观察这种跳变是否在多个关节上同步发生。若多个关节在同一索引附近同时突变，则更可能对应整机构型切换，而不是单一关节的正常局部调整。")
    add_body(doc, "五个纵向站位各自对应一圈闭合喷涂轮廓，理论上同一站位内部的关节角应随轮廓参数连续演化，站位之间的过渡也应保持与上一构型邻近。因而，关节曲线的周期性变化可以视为截面循环运动的正常表现，而超过平滑阈值的孤立尖峰或阶跃则必须单独标记。上述判读原则构成图4-1与图4-2之间的证据链：前者定位具体关节状态，后者把相邻变化量直接与25°阈值比较。")
    add_figure(doc, "fig_3_1_ik_waypoint_angles.png", "4-1", "原始逆运动学路径点的六关节角变化", 12.0, "虚线表示 URDF 中读取的关节位置边界；图中横坐标为未时间参数化的 IK 路径点序号。")
    add_body(doc, "由图4-1可见，六个关节角均处于URDF给定的位置边界内，但各关节曲线在截面连接附近出现阶跃式变化，其中第3、4和5关节的变化尤为明显。边界内并不等价于路径连续，因此仅检查关节是否超限不足以证明轨迹可执行。为定量识别这些阶跃，图4-2进一步统计相邻IK路径点的最大关节跨越。")
    add_body(doc, "图4-2进一步显示，原始相邻 IK 路径点的最大跨越达到 165.88°，主要由第4关节产生；第2、3、5和6关节的最大跨越也均超过 120°。这说明“单点 IK 误差小”并不等于“关节路径连续”。验证文件中的 4.09° 是 Ruckig 轨迹采样后的最大相邻步长，不能替代对原始路径点分支跃迁的审查。")
    add_figure(doc, "fig_3_2_joint_continuity.png", "4-2", "原始 IK 路径点的相邻关节跨越", 14.0, "红色虚线为候选筛选时使用的 25° 平滑阈值；当不存在满足阈值的安全候选时，程序仍可能接受跨越较大的解。")
    add_body(doc, "由图4-2可见，多数关节的最大原始跨越明显高于25°候选平滑阈值，说明当前多初值搜索没有在所有站位连接处保持同一逆解分支。该图揭示的是时间参数化前的几何路径问题，不能通过单纯延长运动时间解决。后续改进应在IK路径点层重新选择连续分支，而不是把大角度构型切换直接交给Ruckig插值。")
    add_heading(doc, "4.3 位置优先回退与可达性投影", 2)
    add_body(doc, "当最优候选的位置误差超过 35 mm 时，程序暂时保持上一法兰姿态，仅以目标 TCP 位置构造回退目标，以提高低仰拱区域的可达性。如果误差仍然过大，则沿机器人基座到目标点方向按 0.95～0.45 的比例向内投影，并记录投影距离。当前验证记录未出现投影点，但在索引 278 处发生一次姿态降级，说明该点以位置优先方式通过。")
    add_heading(doc, "4.4 离散连杆净空与碰撞门禁", 2)
    add_body(doc, "碰撞检查并未加载完整三角网格，而是把相邻关节坐标之间的每段连杆均匀采样 17 个点，并依据对应纵向位置处的解析隧道宽度、拱顶和仰拱计算最小净空。其保守净空指标可概括为")
    add_equation(doc, r"c(\boldsymbol q)=\min_{\boldsymbol p\in\mathcal P(\boldsymbol q)}\{r-|x|,\ z-z_{\mathrm{floor}},\ z_{\mathrm{ceiling}}-z\}-m", "4.3")
    add_body(doc, "式（4.3）中，P(q) 为连杆采样点集合，m=5 mm 为安全裕量。对当前验证记录中的全部抽样状态执行该最小值运算后，validation.json 给出的最小净空为 −0.197606 m。代入安全判据 c(q)≥0 可写为")
    add_equation(doc, r"c_{\min}=-0.197606\ \mathrm m=-197.606\ \mathrm{mm}<0", "4.3a")
    add_body(doc, "由式（4.3a）可知，记录中的最不利状态已经越过简化隧道边界约 197.606 mm，因此不能判为无碰撞；同一验证文件还记录了 4 个负净空抽样状态。该结果是当前输出文件的验收依据。需要指出，该方法只提供快速门禁，没有考虑真实连杆半径、末端喷枪体积、关节外壳和隧道三角网格细节，因此后续仍应使用 MoveIt2/FCL 和真实工位模型复核。")

    add_heading(doc, "第5章  Ruckig 时间参数化与动画建立", 1)
    add_heading(doc, "5.1 路径弧长与最小段时长", 2)
    add_body(doc, "完成 285 个关节路径点求解后，程序依据 TCP 目标点的累计弧长建立各段的名义时间。离散路径累计弧长为")
    add_equation(doc, r"s_k=\sum_{i=1}^{k}\left\|\boldsymbol p_i-\boldsymbol p_{i-1}\right\|_2", "5.1")
    add_body(doc, "式（5.1）是离散折线路径弧长的原始计算式。将 make_targets() 生成的 285 个 TCP 目标点逐段求欧氏距离并累加，可得")
    add_equation(doc, r"s_{284}=\sum_{i=1}^{284}\left\|\boldsymbol p_i-\boldsymbol p_{i-1}\right\|_2=13.535599\ \mathrm m", "5.1a")
    add_body(doc, "式（5.1a）给出的 13.535599 m 是规划 TCP 折线的累计长度，与 Python 重建得到的 planned_path_length_m 完全一致。该长度包含五个闭合截面轮廓以及相邻站位之间的连接段，因此明显大于隧道本身 0.45 m 的纵向长度。目标 TCP 速度设为 0.04 m/s。对第 i 段，Ruckig 的 minimum_duration 参数至少满足")
    add_equation(doc, r"\Delta t_i\ge\frac{s_{i+1}-s_i}{v_{\mathrm{tcp,target}}}", "5.2")
    add_body(doc, "以规划路径中最长的相邻目标段为例，其长度为 0.112999 m。由式（5.2）可得该段在 0.04 m/s 名义速度下的最小时长")
    add_equation(doc, r"\Delta t_{227}\ge\frac{0.112999}{0.04}=2.82498\ \mathrm s", "5.2a")
    add_body(doc, "式（5.2a）表示仅从笛卡尔路程和目标速度出发，该段至少需要 2.825 s；若关节速度、加速度或加加速度约束要求更长时间，Ruckig 会继续增加段时长。式（5.2）只是时间下界，不是“实际 TCP 始终等于 0.04 m/s”的等式约束，因此关节空间平滑后仍可能产生笛卡尔速度波动。")
    add_heading(doc, "5.2 速度、加速度和加加速度约束", 2)
    add_body(doc, "每个相邻路径点对均建立一个六自由度 Ruckig 输入，起点和终点加速度设为零，内部路径点速度由相邻段同向差分估计；若某关节在路径点处发生方向反转，则该关节的路径点速度置零。轨迹满足")
    add_equation(doc, r"|\dot q_j|\le\dot q_{j,\max},\quad |\ddot q_j|\le\ddot q_{j,\max},\quad |\dddot q_j|\le\dddot q_{j,\max}", "5.3")
    add_body(doc, "式（5.3）是关节速度、加速度和加加速度的原始约束关系。当前 URDF/YAML 给出的前三关节速度上限为 3.15 rad/s，即 180.48°/s；程序统一设置加速度上限 0.7 rad/s²、加加速度上限 8 rad/s³。对 CSV 时间序列进行数值微分后，最大速度出现在第3关节，其约束利用率为")
    add_equation(doc, r"\eta_v=\frac{102.50}{180.48}=0.568<1", "5.3a")
    add_body(doc, "最大加速度为 40.107°/s²，而 0.7 rad/s² 换算后同样为 40.107°/s²，因此加速度利用率为")
    add_equation(doc, r"\eta_a=\frac{40.107}{40.107}=1.000", "5.3b")
    add_body(doc, "时间序列数值微分得到的最大加加速度约为 451.02°/s³，8 rad/s³ 对应 458.37°/s³，故")
    add_equation(doc, r"\eta_j=\frac{451.02}{458.37}=0.984<1", "5.3c")
    add_body(doc, "由式（5.3a）至式（5.3c）可知，速度和加加速度没有超过设置上限，但加速度已经达到约束边界，安全余量接近于零。每段轨迹按约 0.04 s 的时间间隔采样，删除段首重复点后串接为完整序列，最终得到 9900 个采样点和 390.59 s 的轨迹时长；validation.json 将时间参数化方法记录为 ruckig-segmented。")
    add_heading(doc, "5.3 时间参数化后的关节运动", 2)
    add_body(doc, "图5-1给出了时间参数化后的六关节角。其总体轮廓与图4-1一致，但路径点之间已被密集插值。第3、4、5和6关节仍可观察到构型区段之间的大范围变化，说明重定时主要控制变化速率，并不会从根本上消除错误的逆解分支。")
    add_figure(doc, "fig_4_4_time_joint_angles.png", "5-1", "Ruckig 时间参数化后的六关节角", 12.0)
    add_body(doc, "由图5-1可见，Ruckig把285个离散关节状态扩展为连续时间序列，使每个构型区段内部的曲线更加平滑；但原始路径点之间的大范围姿态变化仍保留在总体轮廓中。该结果说明时间参数化改善的是速度、加速度和加加速度连续性，而不是重新求解几何路径。为判断平滑后的运动是否接近动力学上限，图5-2进一步给出速度时程和约束利用率。")
    add_body(doc, "图5-2a显示各关节速度的峰值集中在截面衔接和构型切换附近；图5-2b显示峰值速度均未达到设置的速度上限，而峰值加速度利用率接近 100%。这与分段 Ruckig 在每段内满足约束、并频繁在路径点处重新建立边界状态的实现方式一致。")
    add_figure(doc, "fig_4_5_joint_velocity_acceleration.png", "5-2", "关节速度时程及峰值约束利用率", 13.5, "速度和加速度由重建的 Ruckig profile 直接读取；柱状图按各关节 YAML/程序约束归一化。")
    add_body(doc, "由图5-2可知，速度峰值虽然仍低于各关节上限，但加速度峰值几乎达到100%利用率，与式（5.3b）的计算结果一致。这意味着当前轨迹在加速度维度缺少明显安全余量，若真实喷枪质量、控制周期或跟踪误差与仿真假设不同，实机响应可能进一步恶化。因此，该图支持“轨迹完成了约束内重定时”，但不能支持“轨迹具有充足动力学裕量”的结论。")
    add_heading(doc, "5.4 多视角动画生成", 2)
    add_body(doc, "动画程序读取时间序列中的关节角，对每个选定帧执行正运动学，依次获得基座、六个关节、法兰和虚拟 TCP 坐标。主动画采用约 360 个均匀时间采样帧；正视、侧视和等轴测动画分别设置观察角；四视图动画在同一画布中同步显示正视、左前、右前和俯视结果。图5-3截取四视图动画前 30 s 内的四个代表帧，用于说明动画并非预制视频，而是由 CSV 中的关节状态逐帧重建。")
    add_figure(doc, "fig_5_2_animation_frames.png", "5-3", "四视角运动动画的代表帧", 13.0, "四幅图来自 variable_invert_motion_4view.gif 的原始帧提取，未使用生成式图像。该 GIF 默认仅展示完整轨迹的前 30 s。")
    add_body(doc, "图5-3从四个观察方向展示同一时刻的机器人构型和TCP位置，可以辅助识别单一视角下被遮挡的关节折叠、末端朝向和隧道边界关系。四个子图之间的姿态同步表明它们来自同一CSV时间索引，而不是分别制作的示意图。需要强调，动画只能提供直观审查，轨迹是否准确和安全仍必须由后续TCP误差与碰撞数据判定。")

    add_heading(doc, "第6章  仿真结果与定量分析", 1)
    add_heading(doc, "6.1 TCP 三维运动轨迹", 2)
    add_body(doc, "图6-1比较了 Ruckig 段内线性插值的规划 TCP 路径与由关节角经正运动学计算得到的实际 TCP。规划路径长度为 13.536 m，而实际 TCP 轨迹长度为 19.038 m。两者的差异主要集中在截面连接和逆解分支变化附近，并表现为偏离正常马蹄形轮廓的散点和外伸轨迹。")
    add_body(doc, "该三维对比采用统一世界坐标系：闭合截面主要分布在横向—竖向平面，纵向坐标随五个站位递增；规划轨迹以连续灰线表示，实际TCP则按时间着色。若重定时过程保持了原有笛卡尔路径，两类轨迹应在各站位轮廓及站位过渡段上基本重合。反之，彩色轨迹脱离灰线并向隧道外侧伸展，即说明关节空间插值改变了末端在笛卡尔空间中的运动路线。")
    add_body(doc, "轨迹长度差还需要与空间分布共同判断。若实际轨迹只是在同一路径上采用了不同采样密度，其累计长度与规划值应基本一致；当前多出的5.503 m却伴随多处横向和竖向外伸，说明差异来自真实几何偏移，而不是采样点数量不同。图6-1保留完整时间色标，可根据颜色连续性判断异常发生在单个站位内部还是站位切换期间，并据此回查CSV中的对应时间索引。")
    add_body(doc, "从喷涂工艺角度看，外伸轨迹会同时改变喷枪与壁面的距离、入射角和局部停留时间。即使机器人关节曲线在视觉上连续，只要实际TCP脱离规划轮廓，涂层覆盖宽度和沉积厚度就不再具有可预测性。因此，三维轨迹图承担的是几何一致性审查，而不是单纯展示机器人运动外观；它与后续速度和误差时程共同构成判定当前轨迹不可直接执行的证据。")
    add_figure(doc, "fig_4_1_tcp_trajectory_3d.png", "6-1", "规划 TCP 与实际 TCP 的三维轨迹对比", 11.8, "灰线为段内规划插值路径；彩色散点为实际 TCP，颜色表示时间。图例位于左上空白区，色条独立置于右侧。")
    add_body(doc, "由图6-1可见，实际TCP在多数轮廓区段能够沿规划马蹄形路径分布，但在若干站位连接处出现明显外伸散点和跨越轨迹。实际轨迹长度比规划轨迹增加约5.503 m，说明关节空间插值并未始终保持TCP沿原定笛卡尔路径运动。该三维对比为后续速度尖峰和误差峰值提供了空间位置依据。")
    add_heading(doc, "6.2 TCP 位置与速度", 2)
    add_body(doc, "图6-2a中的 X、Y 和 Z 坐标呈现五个截面循环及轴向站位递增特征，但在约 60、90、145、220、300 和 375 s 附近出现异常尖峰。根据实际 TCP 坐标的数值微分，平均速度为 0.0486 m/s，中位速度为 0.0426 m/s，95% 分位速度为 0.1467 m/s，最大速度达到 0.3337 m/s。")
    add_body(doc, "位置时程与速度时程必须联合解释。位置曲线中的周期起伏反映喷枪沿闭合截面循环运动，纵向坐标的台阶则对应机器人从一个站位推进至下一个站位；这些属于规划运动的正常结构。真正需要关注的是三个坐标分量在极短时间内同时发生突变，并在速度曲线上形成窄而高的峰值。该对应关系可用于区分正常截面转弯、数值微分噪声和构型跃迁造成的异常运动。")
    add_body(doc, "本报告以0.04 m/s为喷涂阶段的名义目标速度，但该目标并不是对每一时刻TCP速度的严格等式约束。Ruckig约束的是关节速度、加速度和加加速度；若关节路径本身跨越了错误分支，即使所有关节导数均满足上限，末端仍可能产生远高于工艺目标的瞬时速度。因此，图6-2同时保留目标速度虚线、完整时程和峰值统计，避免只用平均值评价喷涂稳定性。")
    add_figure(doc, "fig_4_2_tcp_position_speed.png", "6-2", "实际 TCP 坐标与速度时程", 13.0, "红色虚线为 0.04 m/s 目标速度。速度由 CSV 中实际 TCP 坐标对时间求数值梯度获得。")
    add_body(doc, "由图6-2可见，正常喷涂区段的TCP速度主要分布在0.04 m/s目标值附近，但多个时刻出现0.15 m/s以上的突增，最大值约为目标速度的8.34倍。这些速度尖峰与三维轨迹中的外伸位置相互对应，表明问题不是一般的采样噪声，而是由构型切换或站位衔接引起的显著空间偏移。因而，平均速度接近目标值不能代表整个喷涂过程满足恒速要求。")
    add_heading(doc, "6.3 轨迹误差", 2)
    add_body(doc, "实际 TCP 路径误差定义为同一时间采样点处实际位置与规划插值位置之间的欧氏距离")
    add_equation(doc, r"e_i=\left\|\boldsymbol p_{\mathrm{tcp},i}^{\mathrm{actual}}-\boldsymbol p_{\mathrm{tcp},i}^{\mathrm{planned}}\right\|_2", "6.1")
    add_body(doc, "最大误差出现在 t=62.664 s 附近。该采样点的实际 TCP 坐标和规划插值坐标分别为")
    add_equation(doc, r"\boldsymbol p^{\mathrm{actual}}=\begin{bmatrix}0.664353\\0.399449\\0.107063\end{bmatrix}\ \mathrm m,\qquad \boldsymbol p^{\mathrm{planned}}=\begin{bmatrix}0.352249\\0\\0.163248\end{bmatrix}\ \mathrm m", "6.1a")
    add_body(doc, "将式（6.1a）的两组真实坐标代入式（6.1），可得")
    add_equation(doc, r"e_{\max}=\sqrt{(0.664353-0.352249)^2+(0.399449-0)^2+(0.107063-0.163248)^2}=0.510025\ \mathrm m", "6.1b")
    add_body(doc, "式（6.1b）将最大误差明确复算为 510.025 mm，与 validation.json 中 realized_tcp_path_error_max_m=0.510025 的记录一致。整体误差中位数为 1.16 mm，但均值为 13.82 mm、95%分位数为 70.59 mm，说明误差分布具有显著长尾；共有 1043/9900 个采样点超过 10 mm，占 10.54%。图6-3a显示异常峰值与截面或构型切换具有重复对应关系，图6-3b进一步表明，少量大误差已经足以破坏喷涂连续性和安全性。")
    add_figure(doc, "fig_4_3_tcp_path_error.png", "6-3", "TCP 路径误差时程与经验累积分布", 14.5, "黑色虚线表示 10 mm 判据。横坐标和图例均与曲线数据区分离。")
    add_body(doc, "由图6-3a可见，误差不是均匀增加，而是在约60、90、145、220、300和375 s附近形成重复峰值；这与图6-2中的速度异常时刻一致。图6-3b显示绝大多数采样点误差较小，但分布尾部延伸至510.03 mm，导致均值显著高于中位数。该结果说明仅报告中位误差会掩盖少量但严重的失效点，必须同时给出最大值、分位数和超限点数量。")
    add_heading(doc, "6.4 验证指标综合评价", 2)
    add_body(doc, "图6-4将四项关键指标与程序采用的判据并列。最大 IK 误差 3.00 mm 和时间参数化后最大相邻关节采样步长 4.09°分别低于 10 mm 与 20°判据；但最大实际 TCP 误差远高于 10 mm，最小采样净空为负。因此，当前 status=review_required 的判定与原始数据一致。")
    add_body(doc, "四项指标分别对应不同验证层级：IK误差检验离散目标点能否被数值求解，相邻采样步长检验重定时后关节序列是否局部平滑，实际TCP误差检验末端是否继续跟踪规划路径，最小净空则检验机器人连杆与隧道边界的几何安全性。前两项通过只能说明局部求解和时间离散没有明显数值失稳，不能替代后两项对完整运动过程的工程验收。")
    add_figure(doc, "fig_5_1_validation_metrics.png", "6-4", "关键验证指标与判据对比", 12.2, "绿色表示通过对应标量判据，红色表示未通过；虚线为判据位置。各面板采用独立横坐标量纲。")
    add_body(doc, "图6-4直观显示，单点IK误差和时间采样步长两项指标通过，但实际TCP最大误差和最小净空两项关键安全指标未通过。由此可见，局部数值求解成功并不能替代整条轨迹的笛卡尔跟踪与碰撞验收。为把图中的四项判据与其余仿真规模、状态数据放在同一位置比较，表6.1对核心结果进行完整汇总。")
    add_three_line_table(
        doc,
        "6.1",
        "仿真核心结果汇总",
        ["指标", "结果", "评价"],
        [["纵向站位", "5", "完成变截面离散"], ["IK 路径点", "285", "形成完整截面路径"], ["时序采样点", "9900", "约 25 Hz"], ["轨迹持续时间", "390.59 s", "完成整段动画数据"], ["最大 IK 误差", "3.00 mm", "通过 10 mm 判据"], ["最大相邻采样步长", "4.09°", "通过 20°判据"], ["姿态降级点", "1（索引278）", "需复核喷枪角度"], ["最大实际 TCP 误差", "510.03 mm", "严重超限"], ["实际误差>10 mm", "1043 点（10.54%）", "未通过"], ["最小采样净空", "−197.61 mm", "发生几何侵入"], ["碰撞采样点", "4", "未通过"], ["验证状态", "review_required", "不得直接下发实体机器人"]],
        [5.6, 4.0, 6.4],
    )
    add_body(doc, "由表6.1可知，当前仿真已经生成5个纵向站位、285个IK路径点和9900个时序采样点，说明数据生产链和动画链已经建立。然而，姿态降级点、510.03 mm最大TCP误差、1043个超限点、−197.61 mm最小净空和4个碰撞采样点共同表明轨迹尚未达到实机下发条件。因此，表中“完成”项描述的是仿真流程完成度，而不是工程验收通过。")

    add_heading(doc, "第7章  讨论与改进建议", 1)
    add_heading(doc, "7.1 已经完成的工作", 2)
    add_body(doc, "从文件证据看，当前成果已经完成了带曲线仰拱变截面隧道的参数化建模、FR5 官方 URDF 加载、虚拟 TCP 和喷涂姿态构造、多初值逆运动学、简化碰撞门禁、分段 Ruckig 时间参数化、CSV/JSON 数据落盘和多视角动画生成。该链条能够复现，且每一步均有源程序或数据文件支撑。为了区分“当前文件夹已经证明的内容”和“任务书要求但尚无证据的内容”，表7.1逐项进行对应审查。")
    add_three_line_table(
        doc,
        "7.1",
        "当前文件夹与任务书要求的对应关系",
        ["任务要求", "文件夹证据", "判定"],
        [["FR5 正逆运动学可视化", "URDF、IK CSV、TCP CSV、GIF", "部分完成；逆解和动画证据充分"], ["至少两种隧道截面", "仅带仰拱变截面成果", "本文件夹证据不足"], ["轨迹图及关节曲线", "本报告由 CSV 重绘", "完成"], ["至少三种路径运动", "当前主要为环向闭合截面+站位推进", "证据不足"], ["轨迹精度分析", "误差和验证 JSON", "已完成初步量化"], ["覆盖均匀性分析", "无涂层沉积/搭接率数据", "未完成"], ["流体仿真", "无 CFD/射流结果", "未完成"], ["机械与电气图纸", "无图纸文件", "未完成"], ["Qt 上位机", "无软件或界面证据", "未完成"]],
        [5.0, 6.6, 4.4],
    )
    add_body(doc, "由表7.1可见，FR5运动学可视化、轨迹图和初步精度分析已有代码与数据支撑；但至少两类隧道截面、至少三种路径运动、覆盖均匀性、流体仿真、机械电气图纸和Qt上位机等任务在当前文件夹中仍缺少完整证据。该表的作用是限定本报告结论范围，避免把“计划完成”或“具备实现可能性”误写成“已经完成”。")
    add_heading(doc, "7.2 当前结果不能直接执行的原因", 2)
    add_body(doc, "第一，数值 IK 在部分截面连接处选择了不同构型分支，原始路径点出现 100°以上跃迁。Ruckig 只能在两个关节状态之间生成满足导数约束的轨迹，不能保证中间 TCP 仍沿规划喷涂轮廓运动。第二，程序以关节线段采样点代替真实连杆体积，既可能漏检，也可能因坐标语义不一致产生保守侵入；无论哪种情况，负净空都要求停止实体部署。第三，150 mm TCP 是虚拟假设值，尚未通过真实喷枪 CAD、质量参数和手眼/工具标定确认。第四，规划显示轮廓、数值法向目标和段内线性规划路径不是完全统一的数据对象，增加了误差解释难度。")
    add_heading(doc, "7.3 面向下一版本的技术改进", 2)
    add_body(doc, "针对前述分支跃迁、TCP跟踪超限、碰撞风险和任务覆盖不足等问题，需要按照对安全性和结果可信度的影响确定修正顺序。表7.2将改进措施划分为P0、P1和P2三个优先级，并为每项措施给出可量化的验收标准，以避免后续修改只停留在代码或动画外观层面。")
    add_three_line_table(
        doc,
        "7.2",
        "建议的修正顺序与验收标准",
        ["优先级", "修正措施", "建议验收标准"],
        [["P0", "在 IK 路径点层消除分支跃迁，禁止>20°跨越进入重定时", "原始与采样关节步长均满足阈值"], ["P0", "在重定时后重新计算 TCP，并对误差闭环迭代", "最大 TCP 误差≤10 mm，超限点为0"], ["P0", "使用 MoveIt2/FCL 完整网格和自碰撞矩阵复核", "全轨迹、工具和过渡段均无碰撞"], ["P1", "将真实喷枪 CAD、质量和标定 TCP 替换150 mm假定值", "工具标定残差和装配参数可追溯"], ["P1", "统一显示路径、求解路径和误差基准路径", "三者来自同一 source-data 对象"], ["P1", "截面间采用明确的喷涂关闭过渡轨迹", "过渡段不计入涂层路径且安全可达"], ["P2", "补充标准马蹄形、纵向条带和螺旋路径", "满足至少两种截面、三种路径要求"], ["P2", "引入喷幅、搭接率和沉积模型", "输出覆盖率、厚度均匀性和参数对比"]],
        [2.0, 8.0, 6.0],
    )
    add_body(doc, "由表7.2可知，P0工作首先要求消除原始IK分支跃迁、把最大TCP误差降低至10 mm以内，并使用完整碰撞模型确认全轨迹安全；这些条件未满足前，不应开展实机喷涂。P1工作用于统一工具标定和路径语义，P2工作则扩展截面类型、运动形式和覆盖均匀性评价。该顺序体现了“先解决安全与正确性，再扩展功能与工艺指标”的工程原则。")

    add_heading(doc, "第8章  结论", 1)
    add_body(doc, "本文基于项目文件夹中的实际代码和数据，对 FR5 带仰拱变截面隧道喷涂运动仿真进行了完整复盘。该仿真以解析几何建立 0.45 m 长、5 站位的变截面隧道，使用 50 mm 离壁距离和 150 mm 虚拟 TCP 构造喷涂目标，通过官方 FR5 URDF、带边界多初值逆运动学、简化碰撞门禁及分段 Ruckig 重定时生成 9900 点运动序列，并形成多视角动画。由此可以确认，运动仿真的基本框架、数据接口和可视化流程已经建立。")
    add_body(doc, "定量结果同时表明，该版本尚未达到工程执行条件。虽然最大 IK 误差和时间采样步长分别为 3.00 mm 和 4.09°，但原始 IK 路径存在 165.88° 的构型分支跨越，实际 TCP 最大偏差达到 510.03 mm，1043 个点超过 10 mm，且最小采样净空为 −197.61 mm。因此，动画能够播放只能证明状态序列可被绘制，不能证明轨迹连续、喷涂均匀或碰撞安全。后续必须先完成分支连续性、TCP 跟踪和完整碰撞模型三项 P0 修正，再扩展流体仿真、覆盖均匀性、末端结构和控制系统成果。")

    add_heading(doc, "参考文献", 1)
    refs = [
        "[1] 丁彦玉. 《机电综合实践》《机电系统课程设计》任务书：隧道衬砌喷涂机器人运动学仿真与路径规划设计[Z]. 2026-06-08.",
        "[2] 国家市场监督管理总局, 国家标准化管理委员会. GB/T 7713.1—2025 信息与文献 编写规则 第1部分：学位论文[S]. 2025.",
        "[3] FAIRINO. FAIRINO-FR5 collaborative robot specifications[EB/OL]. https://www.fairino.com/FR/4.html.",
        "[4] Berscheid L, Kröger T. Jerk-limited real-time trajectory generation with arbitrary target states[C]//Robotics: Science and Systems XVII. 2021. DOI:10.15607/RSS.2021.XVII.015.",
        "[5] SciPy Developers. scipy.optimize.least_squares: bounded nonlinear least-squares[EB/OL]. https://docs.scipy.org/doc/scipy/reference/generated/scipy.optimize.least_squares.html.",
        "[6] Springer Nature. Nature Research Figure Guide: Preparing figures—our specifications[EB/OL]. https://research-figure-guide.nature.com/figures/preparing-figures-our-specifications/.",
        "[7] FAIRINO project assets. fairino5_v6.urdf and fairino5_v6_moveit2_config joint_limits.yaml[CP]. Local project source tree, accessed 2026-07-14.",
        "[8] 本项目源程序. run_fr5_invert_variable_section.py, robot_model.py, time_parameterization.py[CP]. 2026.",
    ]
    for ref in refs:
        p = doc.add_paragraph(style="参考文献")
        set_paragraph_format(p, first_line=False, align=WD_ALIGN_PARAGRAPH.LEFT, line=1.25, after=4)
        p.paragraph_format.left_indent = Pt(24)
        p.paragraph_format.first_line_indent = Pt(-24)
        set_run_font(p.add_run(ref), size=10.5)

    add_heading(doc, "附录A  仿真复现与输出说明", 1)
    add_heading(doc, "A.1 主仿真命令", 2)
    p = doc.add_paragraph()
    set_paragraph_format(p, first_line=False, align=WD_ALIGN_PARAGRAPH.LEFT, line=1.0, before=4, after=4)
    run = p.add_run("python scripts/run_fr5_invert_variable_section.py --out C:\\Users\\86198\\Desktop\\ww --length 0.45 --stations 5 --samples 24 --stand-off 0.05")
    set_run_font(run, size=9.5, cn="Consolas", en="Consolas")
    add_body(doc, "上述命令会生成 IK 路径点、时间参数化 TCP/关节轨迹、验证 JSON、几何图和主动画。多视角动画由 render_variable_invert_multiview.py、render_variable_invert_animation.py 和 render_variable_invert_4view.py 读取同一 TCP CSV 后生成。")
    add_heading(doc, "A.2 结果字段解释", 2)
    add_body(doc, "validation.json是主仿真程序输出的综合验证记录，其中不同字段分别对应逆解精度、TCP跟踪、姿态降级、连杆净空、碰撞计数和时间参数化状态。为便于复现实验后快速判断结果是否发生变化，表A.1列出本文使用的关键字段、物理含义和本次运行数值。")
    add_three_line_table(
        doc,
        "A.1",
        "validation.json 关键字段含义",
        ["字段", "含义", "本次结果"],
        [["ik_error_max", "离散 IK 法兰位置最大误差", "0.002997 m"], ["realized_tcp_path_error_max_m", "实际 TCP 对规划插值路径的最大误差", "0.510025 m"], ["degraded_orientation_points", "姿态降级点索引", "[278]"], ["minimum_sampled_link_clearance_m", "离散连杆采样最小净空", "−0.197606 m"], ["collision_sample_count", "抽样状态中负净空数量", "4"], ["time_parameterization", "时间参数化方法", "ruckig-segmented"], ["status", "综合验证状态", "review_required"]],
        [5.4, 7.0, 3.6],
    )
    add_body(doc, "由表A.1可知，ik_error_max仅为2.997 mm，但realized_tcp_path_error_max_m达到0.510025 m，说明法兰目标点的局部逆解精度与重定时后的TCP路径精度并非同一指标。minimum_sampled_link_clearance_m为负且collision_sample_count为4，与status=review_required相互一致。复现者修改参数后应同时核对这些字段，而不能只观察动画是否平滑。")

    add_heading(doc, "附录B  人工智能辅助使用说明", 1)
    add_body(doc, "本报告使用 OpenAI Codex（GPT-5）进行文件结构梳理、源程序逻辑归纳、CSV/JSON 数据统计、Python 科研配图脚本编写、论文文字结构化和 Word 排版。原始机器人模型、运动仿真程序、CSV/JSON 数据和 GIF 动画均来自项目文件夹，不是由生成式模型虚构；报告中的数值图由 Python 直接读取原始数据计算。按本报告撰写与排版工作量估算，AI 工具参与占比约为 40%。作者应在提交前复核全部技术表述、实验结论和参与占比，并对最终提交内容承担责任。")

    OUT_PATH.parent.mkdir(parents=True, exist_ok=True)
    doc.save(OUT_PATH)
    print(OUT_PATH)


if __name__ == "__main__":
    build_report()
